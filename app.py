"""Excel catalog, versioned baskets, and email orders for the Окунев website."""

import io
import json
import logging
import os
import re
import secrets
import smtplib
import sqlite3
import ssl
import unicodedata
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from email.message import EmailMessage
from email.utils import parseaddr
from pathlib import Path
from typing import Any

import xlrd
from dotenv import load_dotenv
from fastapi import FastAPI, Header, HTTPException, Query, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.exception_handlers import request_validation_exception_handler
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, field_validator
from starlette.datastructures import Headers

import account_store

load_dotenv(Path(__file__).with_name(".env"))
logger = logging.getLogger(__name__)
app = FastAPI(title="Окунев — заказы по прайсу", docs_url=None, redoc_url=None)


@app.exception_handler(RequestValidationError)
async def validation_error(request: Request, error: RequestValidationError):
    if request.url.path == "/api/accounts/password":
        # The default validation details include entered values. Password fields
        # must never be reflected in a response, even when validation fails.
        return JSONResponse({"detail": "Укажите текущий пароль и новый пароль длиной от 8 до 128 символов."},
                            status_code=422, headers={"Cache-Control": "no-store"})
    return await request_validation_exception_handler(request, error)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100
HEADER_SCAN_ROWS = 50
CENT = Decimal("0.01")
MAX_ORDER_TOTAL = Decimal("999999999999.99")
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CATALOG_SCHEMA = """
    CREATE TABLE IF NOT EXISTS catalog_meta (
        singleton INTEGER PRIMARY KEY CHECK (singleton = 1),
        version TEXT NOT NULL, filename TEXT NOT NULL, uploaded_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS products (
        id TEXT PRIMARY KEY, position INTEGER NOT NULL, sku TEXT NOT NULL,
        name TEXT NOT NULL, unit TEXT NOT NULL, price TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS catalog_archive (
        version TEXT PRIMARY KEY, filename TEXT NOT NULL, uploaded_at TEXT NOT NULL,
        products_json TEXT NOT NULL
    );
"""


class CatalogUploadGuard:
    """Reject unauthorized or oversized uploads before multipart files are spooled."""

    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope["type"] != "http" or scope.get("path") != "/api/catalog/import" or scope.get("method") != "POST":
            await self.app(scope, receive, send)
            return
        headers = Headers(scope=scope)
        expected = os.environ.get("ADMIN_TOKEN", "")
        if not expected:
            response = JSONResponse({"detail": "Загрузка прайса выключена. Владелец должен настроить пароль ADMIN_TOKEN."}, status_code=503)
        elif not secrets.compare_digest(headers.get("x-admin-token", "").encode(), expected.encode()):
            response = JSONResponse({"detail": "Неверный пароль владельца."}, status_code=401)
        else:
            response = None
        if response:
            await response(scope, receive, send)
            return
        # Multipart boundaries and headers get a bounded allowance above the file size.
        body_limit = MAX_UPLOAD_BYTES + 256 * 1024
        declared_size = headers.get("content-length")
        if declared_size is not None:
            try:
                size = int(declared_size)
            except ValueError:
                await JSONResponse({"detail": "Некорректный размер запроса."}, status_code=400)(scope, receive, send)
                return
            if size < 0 or size > body_limit:
                await JSONResponse({"detail": "Размер прайса не должен превышать 10 МБ."}, status_code=413)(scope, receive, send)
                return
        received_bytes = 0

        async def bounded_receive():
            nonlocal received_bytes
            message = await receive()
            if message["type"] == "http.request":
                received_bytes += len(message.get("body", b""))
                if received_bytes > body_limit:
                    raise HTTPException(413, "Размер прайса не должен превышать 10 МБ.")
            return message

        await self.app(scope, bounded_receive, send)


app.add_middleware(CatalogUploadGuard)


@contextmanager
def database():
    if account_store.postgres_configured():
        try:
            with account_store.postgres_database() as connection:
                connection.execute("BEGIN IMMEDIATE")
                connection.executescript(CATALOG_SCHEMA)
                connection.commit()
                bootstrap_catalog(connection)
                yield connection
        except HTTPException:
            raise
        except Exception as error:
            logger.warning("Persistent catalog unavailable (%s)", type(error).__name__)
            raise HTTPException(503, "Постоянная база каталога временно недоступна.") from None
        return
    directory = Path(os.environ.get("DATA_DIR", str(Path(__file__).with_name("data"))))
    directory.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(directory / "catalog.sqlite3", timeout=20)
    connection.row_factory = sqlite3.Row
    try:
        connection.executescript(CATALOG_SCHEMA)
        bootstrap_catalog(connection)
        yield connection
    finally:
        connection.close()


def bootstrap_catalog(connection):
    """Seed a fresh database once from a trusted bundled catalog snapshot."""
    source = os.environ.get("INITIAL_CATALOG_FILE", "").strip()
    if not source or connection.execute("SELECT 1 FROM catalog_meta WHERE singleton = 1").fetchone():
        return
    path = Path(source)
    if not path.is_absolute():
        path = Path(__file__).parent / path
    try:
        with path.open(encoding="utf-8") as stream:
            snapshot = json.load(stream)
        if not isinstance(snapshot, dict):
            raise ValueError("Catalog snapshot must be an object")
        for field, limit in {"catalog_version": 100, "filename": 200, "uploaded_at": 100}.items():
            value = snapshot.get(field)
            if not isinstance(value, str) or not value or len(value) > limit:
                raise ValueError(f"Invalid snapshot field: {field}")
        products = snapshot.get("products")
        if not isinstance(products, list) or len(products) > MAX_ROWS:
            raise ValueError("Invalid snapshot product list")
        seen = set()
        rows = []
        for index, product in enumerate(products):
            if not isinstance(product, dict):
                raise ValueError("Invalid snapshot product")
            for field, limit in {"id": 100, "sku": 100, "name": 500, "unit": 50, "price": 40}.items():
                value = product.get(field)
                if not isinstance(value, str) or len(value) > limit:
                    raise ValueError(f"Invalid snapshot product field: {field}")
            if not product["id"] or not product["name"] or product["id"] in seen:
                raise ValueError("Missing or duplicate snapshot product ID/name")
            seen.add(product["id"])
            rows.append((product["id"], index, product["sku"], product["name"],
                         product["unit"], parse_price(product["price"])))
    except (OSError, ValueError, TypeError, KeyError) as error:
        raise RuntimeError("Не удалось загрузить начальный прайс из INITIAL_CATALOG_FILE.") from error
    with connection:
        # A concurrent owner import or another process may have populated the DB while we read JSON.
        connection.execute("BEGIN IMMEDIATE")
        if connection.execute("SELECT 1 FROM catalog_meta WHERE singleton = 1").fetchone():
            return
        connection.execute("DELETE FROM products")
        connection.executemany("INSERT INTO products (id, position, sku, name, unit, price) VALUES (?, ?, ?, ?, ?, ?)", rows)
        connection.execute("INSERT INTO catalog_meta VALUES (1, ?, ?, ?)",
                           (snapshot["catalog_version"], snapshot["filename"], snapshot["uploaded_at"]))


def read_catalog(version=None):
    with database() as connection:
        connection.execute("BEGIN")
        meta = connection.execute("SELECT * FROM catalog_meta WHERE singleton = 1").fetchone()
        if version is not None and (meta is None or meta["version"] != version):
            archived = connection.execute("""
                SELECT version, filename, uploaded_at, products_json FROM catalog_archive WHERE version = ?
            """, (version,)).fetchone()
            if archived is None:
                raise HTTPException(404, "Эта версия прайса недоступна.")
            return {"catalog_version": archived["version"], "filename": archived["filename"],
                    "uploaded_at": archived["uploaded_at"], "products": json.loads(archived["products_json"])}
        products = connection.execute("SELECT id, sku, name, unit, price FROM products ORDER BY position").fetchall()
    return {
        "catalog_version": meta["version"] if meta else None,
        "filename": meta["filename"] if meta else None,
        "uploaded_at": meta["uploaded_at"] if meta else None,
        "products": [dict(product) for product in products],
    }


def clean_text(value: Any, limit=500):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ILLEGAL_CHARACTERS_RE.sub("", str(value)).strip()[:limit]


def normalize_header(value):
    text = unicodedata.normalize("NFKC", clean_text(value)).lower().replace("ё", "е")
    return re.sub(r"[^\w]+", " ", text).strip()


HEADER_ALIASES = {
    "name": {"наименование", "наименование товара", "наименование продукции", "название",
             "название товара", "товар", "номенклатура", "product", "product name", "name", "description"},
    "price": {"цена", "цена руб", "цена рубли", "цена руб шт", "цена с ндс", "цена за ед",
              "цена за единицу", "цена продажи", "розничная цена", "оптовая цена", "стоимость",
              "price", "unit price", "retail price"},
    "sku": {"артикул", "код", "код товара", "код номенклатуры", "sku", "article", "product code", "code"},
    "unit": {"ед изм", "ед измерения", "единица", "единица измерения", "ед", "unit", "units", "uom"},
}


def detect_headers(row):
    mapping = {}
    for index, value in enumerate(row):
        header = normalize_header(value)
        for field, aliases in HEADER_ALIASES.items():
            if field not in mapping and header in aliases:
                mapping[field] = index
    return mapping if "name" in mapping and "price" in mapping else None


def parse_price(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("not a price")
    text = str(value).strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    text = re.sub(r"(?:руб\.?|р\.|₽|RUB)$", "", text, flags=re.IGNORECASE).strip()
    if len(text) > 40:
        raise ValueError("price too long")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        value = Decimal(text)
        if not value.is_finite() or value < 0 or value > Decimal("999999999"):
            raise ValueError("price outside allowed range")
        return str(value.quantize(CENT, rounding=ROUND_HALF_UP))
    except InvalidOperation as error:
        raise ValueError("not a price") from error


def parse_sheet(rows):
    headers = None
    products = []
    skipped = 0
    for index, row in enumerate(rows):
        if index >= MAX_ROWS + HEADER_SCAN_ROWS:
            raise HTTPException(400, "В прайсе слишком много строк. Допустимо до 20 000 позиций.")
        if headers is None:
            headers = detect_headers(row)
            if headers is None and index >= HEADER_SCAN_ROWS - 1:
                return None
            continue
        if not any(value not in (None, "") for value in row):
            continue
        def value(field):
            position = headers.get(field)
            return row[position] if position is not None and position < len(row) else None
        name = clean_text(value("name"))
        try:
            price = parse_price(value("price"))
        except ValueError:
            skipped += 1
            continue
        if not name:
            skipped += 1
            continue
        products.append({"sku": clean_text(value("sku"), 100), "name": name,
                         "unit": clean_text(value("unit"), 50), "price": price})
        if len(products) > MAX_ROWS:
            raise HTTPException(400, "В прайсе слишком много позиций. Допустимо до 20 000.")
    return (products, skipped) if headers else None


def parse_workbook(content, extension):
    try:
        if extension == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 80 * 1024 * 1024:
                    raise HTTPException(400, "Excel-файл слишком велик после распаковки.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook.worksheets:
                    # Supplier files sometimes contain misleading dimensions due to formatting.
                    sheet.reset_dimensions()
                    parsed = parse_sheet(sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True))
                    if parsed and parsed[0]:
                        return parsed
            finally:
                workbook.close()
        else:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                for sheet in workbook.sheets():
                    parsed = parse_sheet(sheet.row_values(index, end_colx=min(sheet.ncols, MAX_COLUMNS))
                                         for index in range(sheet.nrows))
                    if parsed and parsed[0]:
                        return parsed
            finally:
                workbook.release_resources()
    except HTTPException:
        raise
    except Exception as error:
        logger.info("Invalid Excel upload (%s)", type(error).__name__)
        raise HTTPException(400, "Не удалось открыть Excel-файл. Проверьте формат .xlsx или .xls.") from error
    raise HTTPException(400, "Не найдены товары с заголовками «Наименование» и «Цена». Проверьте столбцы прайса.")


def archive_current_catalog(connection):
    meta = connection.execute("SELECT version, filename, uploaded_at FROM catalog_meta WHERE singleton = 1").fetchone()
    if meta is None:
        return
    products = connection.execute("SELECT id, sku, name, unit, price FROM products ORDER BY position").fetchall()
    connection.execute("""
        INSERT INTO catalog_archive (version, filename, uploaded_at, products_json) VALUES (?, ?, ?, ?)
        ON CONFLICT (version) DO NOTHING
    """, (meta["version"], meta["filename"], meta["uploaded_at"],
          json.dumps([dict(product) for product in products], ensure_ascii=False)))


def save_catalog(content, extension, filename):
    products, skipped = parse_workbook(content, extension)
    version = str(uuid.uuid4())
    uploaded_at = datetime.now(timezone.utc).isoformat()
    for index, product in enumerate(products):
        product["id"] = f"{version}:{index + 1}"
    with database() as connection:
        connection.execute("BEGIN IMMEDIATE")
        archive_current_catalog(connection)
        connection.execute("DELETE FROM products")
        connection.executemany("INSERT INTO products (id, position, sku, name, unit, price) VALUES (?, ?, ?, ?, ?, ?)",
                               [(p["id"], index, p["sku"], p["name"], p["unit"], p["price"])
                                for index, p in enumerate(products)])
        connection.execute("""INSERT INTO catalog_meta VALUES (1, ?, ?, ?)
            ON CONFLICT (singleton) DO UPDATE SET version = excluded.version,
            filename = excluded.filename, uploaded_at = excluded.uploaded_at""", (version, filename, uploaded_at))
        connection.commit()
    return {"catalog_version": version, "filename": filename, "uploaded_at": uploaded_at,
            "products": products, "imported_count": len(products), "skipped_count": skipped}


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/api/catalog")
def catalog():
    return read_catalog()


@app.get("/api/catalog/versions/{version}")
def catalog_version_snapshot(version: str):
    return read_catalog(version)


@app.post("/api/catalog/import")
async def import_catalog(file: UploadFile, x_admin_token: str | None = Header(default=None)):
    expected = os.environ.get("ADMIN_TOKEN", "")
    if not expected:
        raise HTTPException(503, "Загрузка прайса выключена. Владелец должен настроить пароль ADMIN_TOKEN.")
    if not secrets.compare_digest((x_admin_token or "").encode(), expected.encode()):
        raise HTTPException(401, "Неверный пароль владельца.")
    filename = clean_text((file.filename or "").replace("\\", "/").split("/")[-1], 200)
    extension = Path(filename).suffix.lower()
    if extension not in {".xlsx", ".xls"}:
        raise HTTPException(400, "Выберите Excel-файл с расширением .xlsx или .xls.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Размер прайса не должен превышать 10 МБ.")
    if not content:
        raise HTTPException(400, "Excel-файл пуст.")
    return await run_in_threadpool(save_catalog, content, extension, filename)


class Customer(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(default="", max_length=160)
    contact: str = Field(default="", max_length=200)
    comment: str = Field(default="", max_length=2000)

    @field_validator("name", "contact")
    @classmethod
    def trim_optional(cls, value):
        return value.strip()


class OrderItem(BaseModel):
    model_config = ConfigDict(extra="forbid")
    product_id: str = Field(min_length=1, max_length=100)
    quantity: Any

    @field_validator("quantity")
    @classmethod
    def validate_quantity(cls, value):
        if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
            raise ValueError("Количество должно быть целым числом.")
        text = str(value).strip().replace(",", ".")
        if len(text) > 40:
            raise ValueError("Количество слишком велико.")
        try:
            quantity = Decimal(text)
            if not quantity.is_finite() or quantity <= 0 or quantity > Decimal("999999"):
                raise ValueError("Количество должно быть больше нуля и не больше 999999.")
            if quantity != quantity.to_integral_value():
                raise ValueError("Количество должно быть целым числом: 1, 2, 3 и так далее.")
        except InvalidOperation as error:
            raise ValueError("Некорректное количество.") from error
        return int(quantity)


class OrderRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    catalog_version: str = Field(min_length=1, max_length=100)
    customer: Customer
    items: list[OrderItem] = Field(min_length=1, max_length=500)


def prepare_order(order):
    catalog = read_catalog()
    if order.catalog_version != catalog["catalog_version"]:
        raise HTTPException(409, "Прайс обновился. Обновите каталог, проверьте цены и наличие в корзине и повторите оформление.")
    products = {product["id"]: product for product in catalog["products"]}
    selected = []
    seen = set()
    total = Decimal("0")
    for item in order.items:
        if item.product_id in seen:
            raise HTTPException(422, "Товар повторяется в заказе. Укажите общее количество одной строкой.")
        if item.product_id not in products:
            raise HTTPException(422, "В заказе есть товар, которого нет в текущем прайсе.")
        seen.add(item.product_id)
        product = products[item.product_id]
        line_total = (Decimal(product["price"]) * item.quantity).quantize(CENT, rounding=ROUND_HALF_UP)
        total += line_total
        if total > MAX_ORDER_TOTAL:
            raise HTTPException(422, "Сумма заказа слишком велика. Максимум — 999 999 999 999,99 ₽.")
        selected.append({**product, "quantity": item.quantity, "line_total": line_total})
    return selected, total


def build_order_excel(order, selected, total, order_id, created_at):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказ"

    def append(target, values):
        target.append(values)
        for cell in target[target.max_row]:
            # Force text: supplier/customer strings beginning with '=' stay literal.
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"

    # Preserve the confirmed supplier template's five columns and their order.
    append(sheet, [created_at.replace(tzinfo=None)])
    sheet["A1"].number_format = "dd.mm.yyyy"
    header_row = 2
    append(sheet, ["Наименование", "Заказ↓", "Ед. изм.", "Цена", "Сумма"])
    for product in selected:
        append(sheet, [product["name"], product["quantity"], product["unit"],
                       float(Decimal(product["price"])), float(product["line_total"])])
    last_product_row = sheet.max_row
    append(sheet, ["Итого", "", "", "", float(total)])
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A5C4C")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[header_row].height = 25
    sheet.freeze_panes = f"B{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:E{last_product_row}"
    for column, width in {"A": 62, "B": 16, "C": 13, "D": 18, "E": 20}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].number_format = "0"
        row[3].number_format = row[4].number_format = '#,##0.00'
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    information = workbook.create_sheet("Информация")
    for values in [["Заказ", order_id], ["Дата (UTC)", created_at.strftime("%d.%m.%Y %H:%M")],
                   ["Покупатель", order.customer.name], ["Контакт", order.customer.contact],
                   ["Комментарий", order.customer.comment]]:
        append(information, values)
    information.column_dimensions["A"].width = 22
    information.column_dimensions["B"].width = 65
    for row in information:
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    information.row_dimensions[5].height = 50
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def order_document(order):
    selected, total = prepare_order(order)
    created_at = datetime.now(timezone.utc)
    order_id = created_at.strftime("%Y%m%d-%H%M%S") + "-" + secrets.token_hex(8)
    filename = f"order-{created_at.strftime('%Y%m%d-%H%M%S')}.xlsx"
    document = build_order_excel(order, selected, total, order_id, created_at)
    return document, filename, order_id, total, selected, created_at


def storage_call(function, *args):
    try:
        return function(*args)
    except account_store.DuplicateAccount:
        raise HTTPException(409, "Аккаунт с этим e-mail уже существует. Войдите с паролем.") from None
    except account_store.InvalidSession:
        raise HTTPException(401, "Сессия завершена. Войдите в аккаунт ещё раз.") from None
    except account_store.IncorrectPassword:
        raise HTTPException(400, "Текущий пароль указан неверно.") from None
    except Exception as error:
        logger.warning("Account storage unavailable (%s)", type(error).__name__)
        raise HTTPException(503, "Постоянная история заказов временно недоступна. Попробуйте позже.") from None


def require_account(request):
    if not account_store.available():
        raise HTTPException(503, "Аккаунты и постоянная история заказов ещё не настроены.")
    user = storage_call(account_store.user_for_session, request.cookies.get("okunev_session"))
    if user is None:
        raise HTTPException(401, "Войдите в аккаунт, чтобы оформить заказ и открыть историю.")
    return user


def checkout_account(request):
    user = require_account(request) if account_store.available() else None
    expected = request.headers.get("x-account-id")
    if user is not None and expected is not None and expected != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    return user


def remember_order(user, order, document, filename, order_id, total, selected, created_at, status):
    if user is None:
        return
    positions = [{"name": product["name"], "unit": product["unit"], "price": product["price"],
                  "quantity": product["quantity"], "line_total": str(product["line_total"])} for product in selected]
    storage_call(account_store.save_order, user["id"], order_id, created_at.isoformat(), filename,
                 status, str(total), positions, order.customer.model_dump(), document)


@app.post("/api/orders/export")
def export_order(order: OrderRequest, request: Request):
    user = checkout_account(request)
    document, filename, order_id, total, selected, created_at = order_document(order)
    remember_order(user, order, document, filename, order_id, total, selected, created_at, "exported")
    return Response(document, media_type=MIME_XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"', "X-Order-Id": order_id,
                             "X-Order-Saved": "true" if user else "false"})


def valid_email(value):
    if "\n" in value or "\r" in value:
        return False
    address = parseaddr(value)[1]
    return bool(address and address.count("@") == 1 and "." in address.split("@")[1] and " " not in address)


def mail_settings():
    host = os.environ.get("SMTP_HOST", "").strip()
    username = os.environ.get("SMTP_USER", "").strip()
    password = os.environ.get("SMTP_PASSWORD", "")
    sender = os.environ.get("SMTP_FROM", "").strip() or username
    recipients = [recipient.strip() for recipient in os.environ.get("ORDERS_TO", "").split(",") if recipient.strip()]
    security = os.environ.get("SMTP_SECURITY", "starttls").strip().lower()
    default_port = "465" if security == "ssl" else "25" if security == "none" else "587"
    try:
        port = int(os.environ.get("SMTP_PORT", default_port))
    except ValueError:
        return None
    if (not host or not sender or not recipients or not valid_email(sender)
            or not all(valid_email(recipient) for recipient in recipients)
            or security not in {"ssl", "starttls", "none"} or not 1 <= port <= 65535
            or bool(username) != bool(password)):
        return None
    return {"host": host, "port": port, "security": security, "username": username,
            "password": password, "sender": sender, "recipients": recipients}


@app.get("/api/config")
def config():
    return {"mail_ready": mail_settings() is not None,
            "history_ready": account_store.available(),
            "upload_requires_password": True, "upload_enabled": bool(os.environ.get("ADMIN_TOKEN", "")),
            "max_upload_mb": MAX_UPLOAD_BYTES // (1024 * 1024)}


@app.post("/api/orders/send")
def send_order(order: OrderRequest, request: Request):
    user = checkout_account(request)
    document, filename, order_id, total, selected, created_at = order_document(order)
    settings = mail_settings()
    remember_order(user, order, document, filename, order_id, total, selected, created_at,
                   "pending" if settings else "failed")
    if settings is None:
        raise HTTPException(503, "Отправка на почту ещё не настроена. Скачайте Excel-заказ или обратитесь к владельцу.")
    message = EmailMessage()
    message["From"] = settings["sender"]
    message["To"] = ", ".join(settings["recipients"])
    message["Subject"] = f"Окунев — заказ {order_id}"
    if valid_email(order.customer.contact) and parseaddr(order.customer.contact)[1] == order.customer.contact:
        message["Reply-To"] = order.customer.contact
    message.set_content(f"Заказ: {order_id}\nПокупатель: {order.customer.name or 'Не указано'}\n"
                        f"Контакт: {order.customer.contact or 'Не указано'}\nКомментарий: {order.customer.comment}\n"
                        f"Итого: {total:.2f} ₽\n\nПозиции и количества заказа — в приложенном Excel-файле.")
    message.add_attachment(document, maintype="application", subtype=MIME_XLSX.split("/", 1)[1], filename=filename)
    smtp_factory = smtplib.SMTP_SSL if settings["security"] == "ssl" else smtplib.SMTP
    kwargs = {"timeout": 20}
    if settings["security"] == "ssl":
        kwargs["context"] = ssl.create_default_context()
    try:
        with smtp_factory(settings["host"], settings["port"], **kwargs) as connection:
            if settings["security"] == "starttls":
                connection.starttls(context=ssl.create_default_context())
            if settings["username"]:
                connection.login(settings["username"], settings["password"])
            refused = connection.send_message(message, from_addr=settings["sender"], to_addrs=settings["recipients"])
            if refused:
                raise smtplib.SMTPRecipientsRefused(refused)
    except (smtplib.SMTPException, OSError, TimeoutError, ssl.SSLError) as error:
        if user:
            storage_call(account_store.update_order_status, user["id"], order_id, "failed")
        logger.warning("Order email failed (%s), order_id=%s", type(error).__name__, order_id)
        raise HTTPException(502, "Не удалось отправить письмо. Заказ сохранён в корзине. Попробуйте позже или скачайте Excel.") from error
    if user:
        storage_call(account_store.update_order_status, user["id"], order_id, "sent")
    return {"ok": True, "order_id": order_id}


class AccountCredentials(BaseModel):
    model_config = ConfigDict(extra="forbid")
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=8, max_length=128)

    @field_validator("email")
    @classmethod
    def normalize_email(cls, value):
        value = value.strip().casefold()
        if not re.fullmatch(r"[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+", value):
            raise ValueError("Укажите корректный e-mail.")
        return value


class AccountProfile(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, str_strip_whitespace=True)
    name: str = Field(max_length=100)
    phone: str = Field(max_length=40)
    company: str = Field(max_length=200)
    delivery_address: str = Field(max_length=500)


class AccountPasswordChange(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    current_password: str = Field(min_length=1, max_length=128)
    new_password: str = Field(min_length=8, max_length=128)


def session_secure(request):
    setting = os.environ.get("SESSION_COOKIE_SECURE", "").strip().lower()
    return setting in {"1", "true", "yes"} if setting else request.url.scheme == "https"


def account_response(request, user, token):
    response = JSONResponse({"available": True, "user": user})
    response.headers["Cache-Control"] = "no-store"
    response.set_cookie("okunev_session", token, max_age=account_store.SESSION_SECONDS,
                        httponly=True, secure=session_secure(request), samesite="lax", path="/")
    return response


@app.get("/api/accounts/me")
def account_me(request: Request):
    available = account_store.available()
    user = storage_call(account_store.user_for_session, request.cookies.get("okunev_session")) if available else None
    return JSONResponse({"available": available, "user": user}, headers={"Cache-Control": "no-store"})


@app.get("/api/accounts/profile")
def account_profile(request: Request, x_account_id: str | None = Header(default=None)):
    user = require_account(request)
    if x_account_id is not None and x_account_id != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    return JSONResponse({"profile": storage_call(account_store.profile, user)},
                        headers={"Cache-Control": "no-store"})


@app.put("/api/accounts/profile")
def update_account_profile(details: AccountProfile, request: Request, x_account_id: str | None = Header(default=None)):
    user = require_account(request)
    if x_account_id is not None and x_account_id != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    return JSONResponse({"profile": storage_call(account_store.save_profile, user, details.model_dump())},
                        headers={"Cache-Control": "no-store"})


@app.post("/api/accounts/register")
def register_account(credentials: AccountCredentials, request: Request):
    if not account_store.available():
        raise HTTPException(503, "Аккаунты и постоянная история заказов ещё не настроены.")
    user, token = storage_call(account_store.register, credentials.email, credentials.password)
    return account_response(request, user, token)


@app.post("/api/accounts/login")
def login_account(credentials: AccountCredentials, request: Request):
    if not account_store.available():
        raise HTTPException(503, "Аккаунты и постоянная история заказов ещё не настроены.")
    result = storage_call(account_store.login, credentials.email, credentials.password)
    if result is None:
        raise HTTPException(401, "Неверный e-mail или пароль.")
    return account_response(request, *result)


@app.post("/api/accounts/logout")
def logout_account(request: Request):
    if account_store.available():
        storage_call(account_store.logout, request.cookies.get("okunev_session"))
    response = JSONResponse({"ok": True}, headers={"Cache-Control": "no-store"})
    response.delete_cookie("okunev_session", path="/", httponly=True,
                           secure=session_secure(request), samesite="lax")
    return response


@app.post("/api/accounts/password")
def update_account_password(details: AccountPasswordChange, request: Request, x_account_id: str | None = Header(default=None)):
    user = require_account(request)
    if x_account_id is not None and x_account_id != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    changed_user, token = storage_call(account_store.change_password, user["id"], request.cookies.get("okunev_session"),
                                      details.current_password, details.new_password)
    return account_response(request, changed_user, token)


@app.get("/api/orders/history")
def order_history(request: Request, offset: int = Query(default=0, ge=0), limit: int = Query(default=50, ge=1, le=100), x_account_id: str | None = Header(default=None)):
    user = require_account(request)
    if x_account_id is not None and x_account_id != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    return JSONResponse(storage_call(account_store.history, user["id"], offset, limit),
                        headers={"Cache-Control": "no-store"})


@app.delete("/api/orders/history/{order_id}")
def delete_order_history(order_id: str, request: Request, x_account_id: str | None = Header(default=None)):
    user = require_account(request)
    if x_account_id is not None and x_account_id != user["id"]:
        raise HTTPException(409, "Аккаунт изменился. Обновите страницу и войдите в нужный аккаунт.")
    if not storage_call(account_store.delete_order, user["id"], order_id):
        raise HTTPException(404, "Заказ не найден.")
    return JSONResponse({"ok": True}, headers={"Cache-Control": "no-store"})


@app.get("/api/orders/history/{order_id}/file")
def order_history_file(order_id: str, request: Request):
    user = require_account(request)
    result = storage_call(account_store.order_file, user["id"], order_id)
    if result is None:
        raise HTTPException(404, "Заказ не найден.")
    filename, document = result
    return Response(document, media_type=MIME_XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})


# API routes take precedence over the frontend's static files.
app.mount("/", StaticFiles(directory=Path(__file__).with_name("web"), html=True, check_dir=False), name="website")
