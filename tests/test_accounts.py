import hashlib
import io
import json
import os
import smtplib
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import app
import account_store
from account_store import postgres_tls_options


PASSWORD = "test-long-password-0123"


def price_file(name="Товар первого прайса", price="12,50"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Прайс"])
    sheet.append(["Наименование", "Заказ↓", "Ед. изм.", "Цена", "Сумма"])
    sheet.append([name, None, "шт", price, None])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class AccountTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.history_dir = str(Path(self.directory.name) / "history")
        self.environment = patch.dict(os.environ, {
            "DATA_DIR": str(Path(self.directory.name) / "catalog"),
            "INITIAL_CATALOG_FILE": "", "ADMIN_TOKEN": "test-owner-password",
            "ORDERS_DATABASE_URL": "", "ORDERS_HISTORY_DIR": self.history_dir,
            "SESSION_COOKIE_SECURE": "0", "SMTP_HOST": "", "SMTP_FROM": "",
            "SMTP_USER": "", "SMTP_PASSWORD": "", "SMTP_PORT": "587",
            "SMTP_SECURITY": "starttls", "ORDERS_TO": "",
        })
        self.environment.start()
        self.client = TestClient(app)
        self.other = TestClient(app)
        self.catalog = self.upload()

    def tearDown(self):
        self.client.close()
        self.other.close()
        self.environment.stop()
        self.directory.cleanup()

    def upload(self, name="Товар первого прайса", price="12,50"):
        response = self.client.post("/api/catalog/import", headers={"X-Admin-Token": "test-owner-password"},
                                    files={"file": ("price.xlsx", price_file(name, price))})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def register(self, client=None, email="buyer-a@example.com"):
        client = client or self.client
        response = client.post("/api/accounts/register", json={"email": email, "password": PASSWORD})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["user"]

    def order(self, quantity=2):
        return {"catalog_version": self.catalog["catalog_version"], "customer": {},
                "items": [{"product_id": self.catalog["products"][0]["id"], "quantity": quantity}]}

    def exported(self, client=None, quantity=2):
        response = (client or self.client).post("/api/orders/export", json=self.order(quantity))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["x-order-saved"], "true")
        return response

    def cart_document(self, quantity=2, include_missing=True):
        product = self.catalog["products"][0]
        items = [{"product_id": product["id"], "quantity": quantity, "product": {
            **product, "available": True, "previous_price": "10.00", "restocked": True}}]
        if include_missing:
            missing = {"id": "previous-catalog:missing", "name": "Ожидаемый товар", "sku": "wait-1", "unit": "шт",
                       "price": "20.00", "available": False, "source_version": "previous-catalog"}
            items.append({"product_id": missing["id"], "quantity": 7, "product": missing})
        return {"catalog_version": self.catalog["catalog_version"],
                "customer": {"name": "Покупатель", "contact": "Контакт", "comment": "Черновик"}, "items": items}

    def save_cart(self, document=None, revision=0, client=None):
        response = (client or self.client).put("/api/accounts/cart", json={
            "revision": revision, "cart": document if document is not None else self.cart_document()})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["cache-control"], "no-store")
        return response.json()

    def shared_order(self, revision=1, quantity=2):
        return {**self.order(quantity), "cart_revision": revision}

    def test_cart_requires_account_and_mismatch_header_never_reads_or_overwrites_other_cart(self):
        document = self.cart_document()
        for method in ("get", "put"):
            kwargs = {"json": {"revision": 0, "cart": document}} if method == "put" else {}
            self.assertEqual(getattr(self.client, method)("/api/accounts/cart", **kwargs).status_code, 401)
            with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": ""}):
                self.assertEqual(getattr(self.client, method)("/api/accounts/cart", **kwargs).status_code, 503)
        user_a = self.register()
        saved_a = self.save_cart(document)
        user_b = self.register(self.other, "buyer-b@example.com")
        self.assertEqual(self.other.get(f"/api/accounts/cart?user_id={user_a['id']}").json(), {"revision": 0, "cart": None})
        for method in ("get", "put"):
            kwargs = {"json": {"revision": 0, "cart": document}} if method == "put" else {}
            response = getattr(self.other, method)("/api/accounts/cart", headers={"X-Account-Id": user_a["id"]}, **kwargs)
            self.assertEqual(response.status_code, 409)
        for injection in ({"user_id": user_a["id"]}, {"email": user_a["email"]}):
            response = self.other.put("/api/accounts/cart", json={"revision": 0, "cart": document, **injection})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved_a)
        self.assertEqual(self.other.get("/api/accounts/me").json()["user"], user_b)

    def test_cart_cross_device_revision_conflict_and_saved_empty_draft_are_authoritative(self):
        user = self.register()
        initial = self.client.get("/api/accounts/cart")
        self.assertEqual(initial.json(), {"revision": 0, "cart": None})
        self.assertEqual(initial.headers["cache-control"], "no-store")
        document = self.cart_document()
        first = self.save_cart(document)
        self.assertEqual(first, {"revision": 1, "cart": document})
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)
        self.assertEqual(self.other.get("/api/accounts/cart").json(), first)
        empty = {**document, "items": []}
        saved_empty = self.save_cart(empty, 1, self.other)
        self.assertEqual(saved_empty, {"revision": 2, "cart": empty})
        stale = self.client.put("/api/accounts/cart", json={"revision": 1, "cart": document})
        self.assertEqual(stale.status_code, 409, stale.text)
        self.assertEqual(stale.headers["cache-control"], "no-store")
        self.assertEqual(stale.json()["code"], "cart_conflict")
        self.assertEqual({key: stale.json()[key] for key in ("revision", "cart")}, saved_empty)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved_empty)

    def test_cart_survives_new_process_and_does_not_change_catalog_profile_or_old_orders(self):
        user = self.register()
        exported = self.exported()
        profile = self.profile(name="Имя", company="Компания")
        self.client.put("/api/accounts/profile", json=profile)
        saved = self.save_cart()
        code = """
import json,sys
from fastapi.testclient import TestClient
from app import app
data=json.load(sys.stdin)
with TestClient(app) as client:
    assert client.post('/api/accounts/login',json={'email':data['email'],'password':data['password']}).status_code==200
    response=client.get('/api/accounts/cart')
    assert response.status_code==200
    print(json.dumps(response.json()))
"""
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                                input=json.dumps({"email": user["email"], "password": PASSWORD}),
                                text=True, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), saved)
        self.assertEqual(self.client.get("/api/accounts/profile").json()["profile"], {"email": user["email"], **profile})
        self.assertEqual(self.client.get("/api/catalog").json()["products"], self.catalog["products"])
        self.assertEqual(self.client.get(f"/api/orders/history/{exported.headers['x-order-id']}/file").content, exported.content)

    def test_cart_validation_rejects_malformed_snapshots_and_preserves_saved_cart(self):
        self.register()
        document = self.cart_document()
        saved = self.save_cart(document)
        invalid = []
        for quantity in (True, 0, -1, 1000000, 1.5, "2", None):
            bad = json.loads(json.dumps(document))
            bad["items"][0]["quantity"] = quantity
            invalid.append(bad)
        for key, value in (("id", "wrong-id"), ("price", "1e3"), ("price", "1,00"),
                           ("price", "1000000000.00"), ("price", "NaN"), ("price", 12.5),
                           ("previous_price", "1.001"), ("name", "x" * 501), ("sku", "x" * 101),
                           ("unit", "x" * 51), ("available", "false"), ("owner_id", "other-user")):
            bad = json.loads(json.dumps(document))
            bad["items"][0]["product"][key] = value
            invalid.append(bad)
        invalid.extend(({**document, "items": document["items"] * 2}, {**document, "user_id": "injected"},
                        {**document, "customer": {"name": "x" * 161}},
                        {**document, "customer": {"contact": "x" * 201}},
                        {**document, "customer": {"comment": "x" * 2001}},
                        {**document, "items": [{**document["items"][0], "product_id": str(index)} for index in range(501)]}))
        for bad in invalid:
            with self.subTest(document=bad):
                self.assertEqual(self.client.put("/api/accounts/cart", json={"revision": 1, "cart": bad}).status_code, 422)
        for revision in (True, -1, 1.5, "1"):
            self.assertEqual(self.client.put("/api/accounts/cart", json={"revision": revision, "cart": document}).status_code, 422)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)

    def test_cart_write_rechecks_revoked_session_inside_transaction(self):
        self.register()
        token = self.client.cookies.get("okunev_session")
        original = account_store.save_cart

        def revoke_then_save(*args):
            account_store.logout(token)
            return original(*args)

        with patch("account_store.save_cart", side_effect=revoke_then_save):
            response = self.client.put("/api/accounts/cart", json={"revision": 0, "cart": self.cart_document()})
        self.assertEqual(response.status_code, 401, response.text)
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM account_carts").fetchone()[0], 0)

    def test_simultaneous_cart_writes_have_one_winner_and_one_explicit_conflict(self):
        user = self.register()
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)
        left, right = self.cart_document(2), self.cart_document(3)
        def save(client, document):
            response = client.put("/api/accounts/cart", json={"revision": 0, "cart": document})
            return response.status_code, response.json()
        with ThreadPoolExecutor(max_workers=2) as executor:
            futures = [executor.submit(save, client, document) for client, document in ((self.client, left), (self.other, right))]
            results = [future.result(timeout=20) for future in futures]
        self.assertEqual(sorted(status for status, _ in results), [200, 409])
        winner = next(body for status, body in results if status == 200)
        conflict = next(body for status, body in results if status == 409)
        self.assertEqual({key: conflict[key] for key in ("revision", "cart")}, winner)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), winner)

    def test_shared_export_atomically_archives_and_clears_only_submitted_quantities(self):
        self.register()
        document = self.cart_document(quantity=5)
        self.save_cart(document)
        response = self.client.post("/api/orders/export", json=self.shared_order(quantity=2))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["x-cart-completed"], "true")
        self.assertEqual(response.headers["x-cart-revision"], "2")
        expected = json.loads(json.dumps(document))
        expected["items"][0]["quantity"] = 3
        self.assertEqual(self.client.get("/api/accounts/cart").json(), {"revision": 2, "cart": expected})
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertEqual(workbook["Заказ"]["B3"].value, 2)
        self.assertEqual(workbook["Заказ"]["D3"].value, 12.5)
        self.assertEqual(self.client.get(f"/api/orders/history/{response.headers['x-order-id']}/file").content, response.content)
        history_before = self.client.get("/api/orders/history").json()
        stale = self.client.post("/api/orders/export", json=self.shared_order(quantity=2))
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(stale.json()["code"], "cart_conflict")
        self.assertEqual(self.client.get("/api/orders/history").json(), history_before)
        self.assertEqual(self.client.get("/api/accounts/cart").json()["cart"], expected)

    def test_shared_export_clears_available_keeps_missing_and_never_trusts_snapshot_prices(self):
        self.register()
        document = self.cart_document()
        document["items"][0]["product"]["price"] = "0.00"
        self.save_cart(document)
        response = self.client.post("/api/orders/export", json=self.shared_order())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), {
            "revision": 2, "cart": {**document, "items": [document["items"][1]]}})
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertEqual(workbook["Заказ"]["D3"].value, 12.5)
        self.assertEqual(workbook["Заказ"].max_row, 4)

    def test_shared_checkout_conflicts_do_not_archive_or_send_and_legacy_orders_do_not_touch_cart(self):
        self.register()
        saved = self.save_cart(quantity_document := self.cart_document(quantity=1))
        for endpoint in ("/api/orders/export", "/api/orders/send"):
            for order in (self.shared_order(0), self.shared_order(1, quantity=2)):
                with patch("app.smtplib.SMTP", autospec=True) as smtp:
                    response = self.client.post(endpoint, json=order)
                self.assertEqual(response.status_code, 409, response.text)
                self.assertEqual(response.json()["code"], "cart_conflict")
                smtp.assert_not_called()
        self.assertEqual(self.client.get("/api/orders/history").json()["total"], 0)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)
        legacy = self.exported()
        self.assertNotIn("x-cart-revision", legacy.headers)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)
        unavailable = json.loads(json.dumps(quantity_document))
        unavailable["items"][0]["product"]["available"] = False
        self.save_cart(unavailable, 1)
        response = self.client.post("/api/orders/export", json=self.shared_order(2, quantity=1))
        self.assertEqual(response.status_code, 409, response.text)

    def test_shared_export_rolls_back_archive_and_cart_when_completion_write_fails(self):
        self.register()
        saved = self.save_cart()
        original = account_store.write_cart
        def fail_after_write(*args):
            original(*args)
            raise OSError("private cart storage failure")
        with patch("account_store.write_cart", side_effect=fail_after_write):
            response = self.client.post("/api/orders/export", json=self.shared_order())
        self.assertEqual(response.status_code, 503, response.text)
        self.assertNotIn("private cart storage failure", response.text)
        self.assertEqual(self.client.get("/api/orders/history").json()["total"], 0)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)

    def test_shared_send_success_clears_available_and_failure_keeps_draft(self):
        self.register()
        saved = self.save_cart()
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", side_effect=smtplib.SMTPException("failure")):
            failed = self.client.post("/api/orders/send", json=self.shared_order())
        self.assertEqual(failed.status_code, 502, failed.text)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            smtp.return_value.__enter__.return_value.send_message.return_value = {}
            sent = self.client.post("/api/orders/send", json=self.shared_order())
        self.assertEqual(sent.status_code, 200, sent.text)
        self.assertEqual(sent.json()["cart_revision"], 2)
        self.assertIs(sent.json()["cart_completed"], True)
        self.assertEqual(self.client.get("/api/accounts/cart").json()["cart"]["items"], [saved["cart"]["items"][1]])
        self.assertEqual({order["status"] for order in self.client.get("/api/orders/history").json()["orders"]}, {"failed", "sent"})

    def test_shared_send_newer_device_draft_is_retained_but_success_is_recorded(self):
        user = self.register()
        self.save_cart()
        self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD})
        newer = self.cart_document(quantity=9)
        def update_during_send(*args, **kwargs):
            self.save_cart(newer, revision=1, client=self.other)
            return {}
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            smtp.return_value.__enter__.return_value.send_message.side_effect = update_during_send
            response = self.client.post("/api/orders/send", json=self.shared_order())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIs(response.json()["cart_completed"], False)
        self.assertEqual(response.json()["cart_revision"], 2)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), {"revision": 2, "cart": newer})
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"][0]["status"], "sent")

    def test_shared_send_unconfigured_mail_never_clears_cart(self):
        self.register()
        saved = self.save_cart()
        response = self.client.post("/api/orders/send", json=self.shared_order())
        self.assertEqual(response.status_code, 503, response.text)
        self.assertEqual(self.client.get("/api/accounts/cart").json(), saved)
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"][0]["status"], "failed")

    def test_disabled_history_is_honest_and_legacy_checkout_remains_anonymous(self):
        with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": ""}):
            self.assertEqual(self.client.get("/api/accounts/me").json(), {"available": False, "user": None})
            self.assertFalse(self.client.get("/api/config").json()["history_ready"])
            credentials = {"email": "buyer@example.com", "password": PASSWORD}
            for route in ("register", "login"):
                self.assertEqual(self.client.post(f"/api/accounts/{route}", json=credentials).status_code, 503)
            self.assertEqual(self.client.get("/api/orders/history").status_code, 503)
            response = self.client.post("/api/orders/export", json=self.order())
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["x-order-saved"], "false")

    def test_configured_history_requires_login_before_checkout(self):
        self.assertTrue(self.client.get("/api/config").json()["history_ready"])
        self.assertEqual(self.client.get("/api/accounts/me").json(), {"available": True, "user": None})
        for endpoint in ("/api/orders/export", "/api/orders/send"):
            self.assertEqual(self.client.post(endpoint, json=self.order()).status_code, 401)
        self.assertEqual(self.client.get("/api/orders/history").status_code, 401)

    def profile(self, **changes):
        return {"name": "", "phone": "", "company": "", "delivery_address": "", **changes}

    def test_profile_requires_login_and_configured_storage(self):
        for method in ("get", "put"):
            kwargs = {"json": self.profile()} if method == "put" else {}
            response = getattr(self.client, method)("/api/accounts/profile", **kwargs)
            self.assertEqual(response.status_code, 401)
            with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": ""}):
                response = getattr(self.client, method)("/api/accounts/profile", **kwargs)
                self.assertEqual(response.status_code, 503)
        self.register()
        self.client.post("/api/accounts/logout")
        self.assertEqual(self.client.get("/api/accounts/profile").status_code, 401)

    def test_existing_account_gets_empty_profile_without_changing_identity_or_history(self):
        user = self.register()
        exported = self.exported()
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            # An existing production account predates the additive profile table.
            connection.execute("DROP TABLE IF EXISTS account_profiles")
        expected = {"email": user["email"], **self.profile()}
        response = self.client.get("/api/accounts/profile")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {"profile": expected})
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user)
        self.assertEqual(self.client.get(f"/api/orders/history/{exported.headers['x-order-id']}/file").content,
                         exported.content)

    def test_profile_save_trims_fields_survives_new_process_and_can_be_cleared(self):
        user = self.register()
        values = self.profile(name="  Александр  ", phone=" +7 (999) 123-45-67 ",
                              company=" ООО Окунев ", delivery_address=" Москва, ул. Лесная, 1 ")
        response = self.client.put("/api/accounts/profile", json=values)
        self.assertEqual(response.status_code, 200, response.text)
        expected = {"email": user["email"], **{key: value.strip() for key, value in values.items()}}
        self.assertEqual(response.json(), {"profile": expected})
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(self.client.get("/api/accounts/profile").json()["profile"], expected)
        code = """
import json,sys
from fastapi.testclient import TestClient
from app import app
data=json.load(sys.stdin)
with TestClient(app) as client:
    login=client.post('/api/accounts/login',json={'email':data['email'],'password':data['password']})
    assert login.status_code==200
    response=client.get('/api/accounts/profile')
    assert response.status_code==200
    print(json.dumps(response.json()))
"""
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                                input=json.dumps({"email": user["email"], "password": PASSWORD}),
                                text=True, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {"profile": expected})
        cleared = self.client.put("/api/accounts/profile", json=self.profile())
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertEqual(cleared.json(), {"profile": {"email": user["email"], **self.profile()}})
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user)

    def test_profile_is_private_and_owner_identity_cannot_be_injected(self):
        user_a = self.register()
        profile_a = self.profile(name="Покупатель А", phone="111", company="Компания А", delivery_address="Адрес А")
        self.assertEqual(self.client.put("/api/accounts/profile", json=profile_a).status_code, 200)
        user_b = self.register(self.other, "buyer-b@example.com")
        initial_b = self.other.get(f"/api/accounts/profile?user_id={user_a['id']}")
        self.assertEqual(initial_b.json(), {"profile": {"email": user_b["email"], **self.profile()}})
        for injected in ({"user_id": user_a["id"]}, {"email": user_a["email"]}, {"id": user_a["id"]}):
            self.assertEqual(self.other.put("/api/accounts/profile", json={**self.profile(), **injected}).status_code, 422)
        profile_b = self.profile(name="Покупатель Б", company="Компания Б")
        self.assertEqual(self.other.put("/api/accounts/profile", json=profile_b).status_code, 200)
        self.assertEqual(self.client.get("/api/accounts/profile").json(),
                         {"profile": {"email": user_a["email"], **profile_a}})
        self.assertEqual(self.other.get("/api/accounts/profile").json(),
                         {"profile": {"email": user_b["email"], **profile_b}})

    def test_profile_validation_rejects_invalid_fields_without_overwriting_saved_values(self):
        self.register()
        saved = self.profile(name="Сохранённое имя", company="Компания")
        self.assertEqual(self.client.put("/api/accounts/profile", json=saved).status_code, 200)
        invalid = []
        for field, limit in (("name", 100), ("phone", 40), ("company", 200), ("delivery_address", 500)):
            invalid.append(self.profile(**{field: "я" * (limit + 1)}))
            for value in (None, 42, True, []):
                invalid.append(self.profile(**{field: value}))
        invalid.extend(({"name": "Only one field"}, {**self.profile(), "unexpected": "value"}))
        for values in invalid:
            with self.subTest(values=values):
                self.assertEqual(self.client.put("/api/accounts/profile", json=values).status_code, 422)
        self.assertEqual(self.client.get("/api/accounts/profile").json()["profile"]["name"], saved["name"])
        boundary = self.profile(name="я" * 100, phone="1" * 40, company="я" * 200, delivery_address="я" * 500)
        self.assertEqual(self.client.put("/api/accounts/profile", json=boundary).status_code, 200)

    def test_stale_profile_form_cannot_read_or_overwrite_after_account_switch(self):
        user_a = self.register()
        profile_a = self.profile(name="Покупатель А")
        self.assertEqual(self.client.put("/api/accounts/profile", json=profile_a).status_code, 200)
        user_b = self.register(self.other, "buyer-b@example.com")
        stale_header = {"X-Account-Id": user_a["id"]}
        self.assertEqual(self.other.get("/api/accounts/profile", headers=stale_header).status_code, 409)
        self.assertEqual(self.other.get("/api/orders/history", headers=stale_header).status_code, 409)
        for endpoint in ("/api/orders/export", "/api/orders/send"):
            self.assertEqual(self.other.post(endpoint, headers=stale_header, json=self.order()).status_code, 409)
        self.assertEqual(self.other.get("/api/orders/history").json()["total"], 0)
        self.assertEqual(self.other.put("/api/accounts/profile", headers=stale_header, json=profile_a).status_code, 409)
        self.assertEqual(self.other.get("/api/accounts/profile").json(),
                         {"profile": {"email": user_b["email"], **self.profile()}})
        self.assertEqual(self.client.get("/api/accounts/profile", headers=stale_header).status_code, 200)

    def test_profile_changes_never_rewrite_historical_customer_or_excel(self):
        self.register()
        order = self.order()
        order["customer"] = {"name": "Первое имя", "contact": "Первый контакт", "comment": "Первый комментарий"}
        exported = self.client.post("/api/orders/export", json=order)
        self.assertEqual(exported.status_code, 200, exported.text)
        original_history = self.client.get("/api/orders/history").json()
        response = self.client.put("/api/accounts/profile", json=self.profile(
            name="Новое имя", phone="Новый телефон", company="Новая компания", delivery_address="Новый адрес"))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get("/api/orders/history").json(), original_history)
        restored = self.client.get(f"/api/orders/history/{exported.headers['x-order-id']}/file")
        self.assertEqual(restored.content, exported.content)

    def test_password_and_session_are_hashed_cookie_is_private_and_email_normalized(self):
        response = self.client.post("/api/accounts/register", json={"email": " Buyer-A@Example.COM ", "password": PASSWORD})
        self.assertEqual(response.status_code, 200, response.text)
        user = response.json()["user"]
        self.assertEqual(user["email"], "buyer-a@example.com")
        self.assertEqual(set(user), {"id", "email"})
        self.assertIn("HttpOnly", response.headers["set-cookie"])
        self.assertIn("SameSite=lax", response.headers["set-cookie"])
        self.assertIn("Max-Age=2592000", response.headers["set-cookie"])
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            stored_password = connection.execute("SELECT password_hash FROM users").fetchone()[0]
            stored_token = connection.execute("SELECT token_hash FROM sessions").fetchone()[0]
        self.assertNotIn(PASSWORD, stored_password)
        self.assertTrue(stored_password.startswith("pbkdf2_sha256$600000$"))
        token = self.client.cookies.get("okunev_session")
        self.assertNotEqual(stored_token, token)
        self.assertEqual(stored_token, hashlib.sha256(token.encode()).hexdigest())
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user)

    def test_secure_cookie_setting_for_https(self):
        with patch.dict(os.environ, {"SESSION_COOKIE_SECURE": "1"}), TestClient(app, base_url="https://testserver") as client:
            response = client.post("/api/accounts/register", json={"email": "secure@example.com", "password": PASSWORD})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("Secure", response.headers["set-cookie"])
            self.assertTrue(client.get("/api/accounts/me").json()["user"])
            rotated = client.post("/api/accounts/password", json={"current_password": PASSWORD,
                                                                   "new_password": "new-secure-password-0123"})
            self.assertEqual(rotated.status_code, 200, rotated.text)
            self.assertIn("Secure", rotated.headers["set-cookie"])
            self.assertTrue(client.get("/api/accounts/me").json()["user"])

    def test_wrong_password_unknown_account_and_duplicate_registration(self):
        self.register()
        duplicate = self.other.post("/api/accounts/register", json={"email": "BUYER-A@EXAMPLE.COM", "password": PASSWORD})
        self.assertEqual(duplicate.status_code, 409)
        responses = []
        for email in ("buyer-a@example.com", "unknown@example.com"):
            response = self.other.post("/api/accounts/login", json={"email": email, "password": "wrong-password"})
            self.assertEqual(response.status_code, 401)
            responses.append(response.json())
        self.assertEqual(responses[0], responses[1])
        self.assertIsNone(self.other.get("/api/accounts/me").json()["user"])
        for credentials in ({"email": "invalid", "password": PASSWORD},
                            {"email": "valid@example.com", "password": "short"}):
            self.assertEqual(self.other.post("/api/accounts/register", json=credentials).status_code, 422)

    def test_password_change_requires_live_account_and_configured_storage(self):
        body = {"current_password": PASSWORD, "new_password": "new-test-password-0123"}
        self.assertEqual(self.client.post("/api/accounts/password", json=body).status_code, 401)
        with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": ""}):
            self.assertEqual(self.client.post("/api/accounts/password", json=body).status_code, 503)
        self.register()
        self.client.post("/api/accounts/logout")
        self.assertEqual(self.client.post("/api/accounts/password", json=body).status_code, 401)

    def account_secrets(self):
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            return (connection.execute("SELECT id, password_hash FROM users ORDER BY id").fetchall(),
                    connection.execute("SELECT token_hash, user_id, expires_at FROM sessions ORDER BY token_hash").fetchall())

    def test_wrong_current_password_does_not_change_password_or_revoke_sessions(self):
        user = self.register()
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)
        before = self.account_secrets()
        token = self.client.cookies.get("okunev_session")
        body = {"current_password": "wrong-current-secret", "new_password": "unused-new-secret-0123"}
        response = self.client.post("/api/accounts/password", json=body)
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(response.json(), {"detail": "Текущий пароль указан неверно."})
        for password in body.values():
            self.assertNotIn(password, response.text)
        self.assertEqual(self.account_secrets(), before)
        self.assertEqual(self.client.cookies.get("okunev_session"), token)
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user)
        self.assertEqual(self.other.get("/api/accounts/me").json()["user"], user)

    def test_password_change_rotates_cookie_revokes_old_devices_and_preserves_private_data(self):
        user_a = self.register()
        self.client.put("/api/accounts/profile", json=self.profile(name="Имя А", company="Компания А"))
        profile = self.client.get("/api/accounts/profile").json()
        exported = self.exported()
        history = self.client.get("/api/orders/history").json()
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user_a["email"], "password": PASSWORD}).status_code, 200)
        old_token = self.client.cookies.get("okunev_session")
        old_device_token = self.other.cookies.get("okunev_session")
        before_hash = self.account_secrets()[0][0][1]
        new_password = "new-test-password-0123"
        with TestClient(app) as buyer_b:
            user_b = self.register(buyer_b, "buyer-b@example.com")
            b_token = buyer_b.cookies.get("okunev_session")
            response = self.client.post("/api/accounts/password", headers={"X-Account-Id": user_a["id"]},
                                        json={"current_password": PASSWORD, "new_password": new_password})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json(), {"available": True, "user": user_a})
            self.assertEqual(response.headers["cache-control"], "no-store")
            self.assertIn("HttpOnly", response.headers["set-cookie"])
            self.assertIn("SameSite=lax", response.headers["set-cookie"])
            new_token = self.client.cookies.get("okunev_session")
            self.assertNotIn(new_token, (old_token, old_device_token))
            self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user_a)
            self.assertIsNone(self.other.get("/api/accounts/me").json()["user"])
            self.assertEqual(self.other.get("/api/accounts/profile").status_code, 401)
            with TestClient(app) as replay:
                replay.cookies.set("okunev_session", old_token)
                self.assertEqual(replay.get("/api/accounts/profile").status_code, 401)
            self.assertEqual(self.other.post("/api/accounts/login", json={"email": user_a["email"], "password": PASSWORD}).status_code, 401)
            self.assertEqual(self.other.post("/api/accounts/login", json={"email": user_a["email"], "password": new_password}).status_code, 200)
            self.assertEqual(buyer_b.cookies.get("okunev_session"), b_token)
            self.assertEqual(buyer_b.get("/api/accounts/me").json()["user"], user_b)
            self.assertEqual(buyer_b.post("/api/accounts/login", json={"email": user_b["email"], "password": PASSWORD}).status_code, 200)
        self.assertEqual(self.client.get("/api/accounts/profile").json(), profile)
        self.assertEqual(self.client.get("/api/orders/history").json(), history)
        self.assertEqual(self.client.get(f"/api/orders/history/{exported.headers['x-order-id']}/file").content, exported.content)
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            stored_hash = connection.execute("SELECT password_hash FROM users WHERE id = ?", (user_a["id"],)).fetchone()[0]
            tokens = {row[0] for row in connection.execute("SELECT token_hash FROM sessions WHERE user_id = ?", (user_a["id"],))}
        self.assertNotEqual(stored_hash, before_hash)
        self.assertTrue(stored_hash.startswith("pbkdf2_sha256$600000$"))
        for password in (PASSWORD, new_password):
            self.assertNotIn(password, stored_hash)
            self.assertNotIn(password, response.text)
        self.assertNotIn(hashlib.sha256(old_token.encode()).hexdigest(), tokens)
        self.assertNotIn(hashlib.sha256(old_device_token.encode()).hexdigest(), tokens)
        self.assertIn(hashlib.sha256(new_token.encode()).hexdigest(), tokens)

    def test_stale_account_header_cannot_change_another_accounts_password(self):
        user_a = self.register()
        self.register(self.other, "buyer-b@example.com")
        before = self.account_secrets()
        response = self.other.post("/api/accounts/password", headers={"X-Account-Id": user_a["id"]},
                                   json={"current_password": PASSWORD, "new_password": "new-test-password-0123"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.account_secrets(), before)

    def test_password_change_validation_rejects_invalid_fields_without_echoing_secrets(self):
        self.register()
        before = self.account_secrets()
        valid = {"current_password": PASSWORD, "new_password": "new-test-password-0123"}
        invalid = [{}, {"current_password": PASSWORD}, {"new_password": valid["new_password"]},
                   {**valid, "email": "injected@example.com"}, {**valid, "user_id": "other-user"},
                   {**valid, "current_password": ""}, {**valid, "current_password": "c" * 129},
                   {**valid, "new_password": "short"}, {**valid, "new_password": "n" * 129}]
        for field in valid:
            invalid.extend({**valid, field: value} for value in (None, 42, True, []))
        for body in invalid:
            with self.subTest(body=body):
                response = self.client.post("/api/accounts/password", json=body)
                self.assertEqual(response.status_code, 422, response.text)
                for password in (body.get("current_password"), body.get("new_password")):
                    if isinstance(password, str) and password:
                        self.assertNotIn(password, response.text)
        self.assertEqual(self.account_secrets(), before)

    def test_password_change_preserves_whitespace_and_accepts_password_length_boundaries(self):
        user = self.register()
        current = PASSWORD
        for new_password in (" 123456 ", "n" * 128):
            response = self.client.post("/api/accounts/password", json={"current_password": current, "new_password": new_password})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": new_password}).status_code, 200)
            if new_password.strip() != new_password:
                self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": new_password.strip()}).status_code, 422)
            current = new_password

    def test_password_change_rechecks_revocation_or_expiry_inside_transaction(self):
        user = self.register()
        original_change = account_store.change_password
        for condition in ("revoked", "expired"):
            with self.subTest(condition=condition):
                self.assertEqual(self.client.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)
                def invalidate_before_transaction(*args):
                    token = args[1]
                    if condition == "revoked":
                        account_store.logout(token)
                    else:
                        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
                            connection.execute("UPDATE sessions SET expires_at = ? WHERE token_hash = ?",
                                               (int(time.time()) - 1, hashlib.sha256(token.encode()).hexdigest()))
                    return original_change(*args)
                with patch("account_store.change_password", side_effect=invalidate_before_transaction):
                    response = self.client.post("/api/accounts/password", json={"current_password": PASSWORD, "new_password": "new-test-password-0123"})
                self.assertEqual(response.status_code, 401, response.text)
                self.assertNotIn("set-cookie", response.headers)
                self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)

    def test_concurrent_password_changes_allow_only_one_rotation_from_same_old_cookie(self):
        user = self.register()
        token = self.client.cookies.get("okunev_session")
        def change(new_password):
            with TestClient(app) as client:
                client.cookies.set("okunev_session", token, domain="testserver.local", path="/")
                response = client.post("/api/accounts/password", json={"current_password": PASSWORD, "new_password": new_password})
                return response.status_code, new_password, client.cookies.get("okunev_session")
        with ThreadPoolExecutor(max_workers=2) as workers:
            results = list(workers.map(change, ("first-new-password-0123", "second-new-password-0123")))
        self.assertEqual(sorted(result[0] for result in results), [200, 401])
        winner = next(result for result in results if result[0] == 200)
        loser = next(result for result in results if result[0] != 200)
        self.assertNotEqual(winner[2], token)
        self.assertIsNone(self.client.get("/api/accounts/me").json()["user"])
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": loser[1]}).status_code, 401)
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": winner[1]}).status_code, 200)

    def test_password_change_rolls_back_password_and_revocation_if_new_session_write_fails(self):
        user = self.register()
        self.assertEqual(self.other.post("/api/accounts/login", json={"email": user["email"], "password": PASSWORD}).status_code, 200)
        before = self.account_secrets()
        with patch("account_store.new_session", side_effect=OSError("private storage failure")):
            response = self.client.post("/api/accounts/password", json={"current_password": PASSWORD, "new_password": "new-test-password-0123"})
        self.assertEqual(response.status_code, 503, response.text)
        self.assertNotIn("private storage failure", response.text)
        self.assertNotIn("set-cookie", response.headers)
        self.assertEqual(self.account_secrets(), before)
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user)
        self.assertEqual(self.other.get("/api/accounts/me").json()["user"], user)

    def test_accounts_cannot_list_or_download_each_others_orders(self):
        user_a = self.register()
        exported_a = self.exported()
        id_a = exported_a.headers["x-order-id"]
        self.register(self.other, "buyer-b@example.com")
        self.assertEqual(self.other.get("/api/orders/history").json()["orders"], [])
        self.assertEqual(self.other.get(f"/api/orders/history/{id_a}/file").status_code, 404)
        exported_b = self.exported(self.other, 3)
        id_b = exported_b.headers["x-order-id"]
        history_a = self.client.get("/api/orders/history").json()
        history_b = self.other.get(f"/api/orders/history?user_id={user_a['id']}").json()
        self.assertEqual([order["id"] for order in history_a["orders"]], [id_a])
        self.assertEqual([order["id"] for order in history_b["orders"]], [id_b])
        self.assertEqual(history_a["orders"][0]["positions"], [{"name": "Товар первого прайса", "unit": "шт",
                          "price": "12.50", "quantity": 2, "line_total": "25.00"}])
        self.assertEqual(history_a["orders"][0]["status"], "exported")
        self.assertEqual(history_a["orders"][0]["customer"], {"name": "", "contact": "", "comment": ""})
        self.assertEqual(self.client.get(f"/api/orders/history/{id_b}/file").status_code, 404)
        self.assertEqual(self.client.get(f"/api/orders/history/{id_a}/file").content, exported_a.content)
        injected = self.order()
        injected["user_id"] = user_a["id"]
        self.assertEqual(self.other.post("/api/orders/export", json=injected).status_code, 422)

    def test_history_and_original_excel_survive_new_process_and_login(self):
        user = self.register()
        exported = self.exported()
        order_id = exported.headers["x-order-id"]
        code = """
import hashlib,json,sys
from fastapi.testclient import TestClient
from app import app
data=json.load(sys.stdin)
with TestClient(app) as client:
    login=client.post('/api/accounts/login',json={'email':data['email'],'password':data['password']})
    assert login.status_code==200
    history=client.get('/api/orders/history').json()
    document=client.get('/api/orders/history/'+data['order_id']+'/file')
    assert document.status_code==200
    print(json.dumps({'user':login.json()['user'],'ids':[order['id'] for order in history['orders']],
                      'file_hash':hashlib.sha256(document.content).hexdigest()}))
"""
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                                input=json.dumps({"email": user["email"], "password": PASSWORD, "order_id": order_id}),
                                text=True, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        restarted = json.loads(result.stdout)
        self.assertEqual(restarted["user"], user)
        self.assertEqual(restarted["ids"], [order_id])
        self.assertEqual(restarted["file_hash"], hashlib.sha256(exported.content).hexdigest())

    def test_replacing_catalog_never_changes_old_history_or_file(self):
        self.register()
        exported = self.exported()
        order_id = exported.headers["x-order-id"]
        old_order = self.order()
        self.upload("Совсем новый товар", "99,99")
        self.assertEqual(self.client.post("/api/orders/export", json=old_order).status_code, 409)
        restored = self.client.get(f"/api/orders/history/{order_id}/file")
        self.assertEqual(restored.content, exported.content)
        workbook = load_workbook(io.BytesIO(restored.content), data_only=True)
        self.assertEqual(workbook["Заказ"]["A3"].value, "Товар первого прайса")
        self.assertEqual(workbook["Заказ"]["D3"].value, 12.5)
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"][0]["total"], "25.00")

    def test_logout_revokes_copied_token_and_expired_sessions_fail(self):
        self.register()
        token = self.client.cookies.get("okunev_session")
        self.assertEqual(self.client.post("/api/accounts/logout").status_code, 200)
        self.other.cookies.set("okunev_session", token)
        self.assertIsNone(self.other.get("/api/accounts/me").json()["user"])
        self.assertEqual(self.other.get("/api/orders/history").status_code, 401)
        response = self.client.post("/api/accounts/login", json={"email": "buyer-a@example.com", "password": PASSWORD})
        self.assertEqual(response.status_code, 200)
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            connection.execute("UPDATE sessions SET expires_at = ?", (int(time.time()) - 1,))
        self.assertIsNone(self.client.get("/api/accounts/me").json()["user"])
        self.assertEqual(self.client.get("/api/orders/history").status_code, 401)

    def test_history_pagination_is_bounded_and_complete(self):
        self.register()
        ids = {self.exported(quantity=quantity).headers["x-order-id"] for quantity in (1, 2, 3)}
        first = self.client.get("/api/orders/history?limit=2").json()
        second = self.client.get("/api/orders/history?limit=2&offset=2").json()
        self.assertEqual(first["total"], 3)
        self.assertEqual(first["next_offset"], 2)
        self.assertIsNone(second["next_offset"])
        self.assertEqual({order["id"] for order in first["orders"] + second["orders"]}, ids)
        self.assertEqual(self.client.get("/api/orders/history?limit=101").status_code, 422)
        self.assertEqual(self.client.get("/api/orders/history?offset=-1").status_code, 422)

    def test_delete_saved_order_requires_login_and_configured_storage(self):
        endpoint = "/api/orders/history/unknown-order"
        self.assertEqual(self.client.delete(endpoint).status_code, 401)
        with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": ""}):
            self.assertEqual(self.client.delete(endpoint).status_code, 503)
        self.register()
        self.client.post("/api/accounts/logout")
        self.assertEqual(self.client.delete(endpoint).status_code, 401)

    def test_nonowner_cannot_delete_saved_order_and_missing_order_is_indistinguishable(self):
        user_a = self.register()
        exported_a = self.exported()
        order_a = exported_a.headers["x-order-id"]
        user_b = self.register(self.other, "buyer-b@example.com")
        exported_b = self.exported(self.other, 3)
        history_a = self.client.get("/api/orders/history").json()
        history_b = self.other.get("/api/orders/history").json()
        forbidden = self.other.delete(f"/api/orders/history/{order_a}?user_id={user_a['id']}",
                                      headers={"X-Account-Id": user_b["id"]})
        missing = self.other.delete("/api/orders/history/unknown-order")
        self.assertEqual(forbidden.status_code, 404)
        self.assertEqual(missing.status_code, 404)
        self.assertEqual(forbidden.json(), missing.json())
        self.assertEqual(self.client.get("/api/orders/history").json(), history_a)
        self.assertEqual(self.other.get("/api/orders/history").json(), history_b)
        self.assertEqual(self.client.get(f"/api/orders/history/{order_a}/file").content, exported_a.content)
        self.assertEqual(self.other.get(f"/api/orders/history/{exported_b.headers['x-order-id']}/file").content,
                         exported_b.content)

    def test_delete_saved_order_removes_only_selected_archive_and_updates_pagination(self):
        user_a = self.register()
        profile_a = self.profile(name="Имя покупателя", company="Компания")
        self.client.put("/api/accounts/profile", json=profile_a)
        exported_a = [self.exported(quantity=quantity) for quantity in (2, 3, 4)]
        removed_id = exported_a[1].headers["x-order-id"]
        remaining_ids = {exported.headers["x-order-id"] for exported in (exported_a[0], exported_a[2])}
        self.register(self.other, "buyer-b@example.com")
        exported_b = self.exported(self.other, 5)
        history_b = self.other.get("/api/orders/history").json()
        catalog = self.client.get("/api/catalog").json()
        response = self.client.delete(f"/api/orders/history/{removed_id}", headers={"X-Account-Id": user_a["id"]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), {"ok": True})
        self.assertEqual(response.headers["cache-control"], "no-store")
        first = self.client.get("/api/orders/history?limit=1").json()
        second = self.client.get("/api/orders/history?limit=1&offset=1").json()
        self.assertEqual(first["total"], 2)
        self.assertEqual(first["next_offset"], 1)
        self.assertIsNone(second["next_offset"])
        self.assertEqual({first["orders"][0]["id"], second["orders"][0]["id"]}, remaining_ids)
        self.assertEqual(self.client.get(f"/api/orders/history/{removed_id}/file").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/orders/history/{removed_id}").status_code, 404)
        with sqlite3.connect(Path(self.history_dir) / "accounts.sqlite3") as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM orders WHERE id = ?", (removed_id,)).fetchone()[0], 0)
        for exported in (exported_a[0], exported_a[2]):
            self.assertEqual(self.client.get(f"/api/orders/history/{exported.headers['x-order-id']}/file").content,
                             exported.content)
        self.assertEqual(self.other.get("/api/orders/history").json(), history_b)
        self.assertEqual(self.other.get(f"/api/orders/history/{exported_b.headers['x-order-id']}/file").content,
                         exported_b.content)
        self.assertEqual(self.client.get("/api/catalog").json(), catalog)
        self.assertEqual(self.client.get("/api/accounts/me").json()["user"], user_a)
        self.assertEqual(self.client.get("/api/accounts/profile").json()["profile"],
                         {"email": user_a["email"], **profile_a})

    def test_stale_account_precondition_cannot_delete_new_accounts_own_saved_order(self):
        user_a = self.register()
        self.register(self.other, "buyer-b@example.com")
        exported_b = self.exported(self.other)
        order_b = exported_b.headers["x-order-id"]
        stale = self.other.delete(f"/api/orders/history/{order_b}", headers={"X-Account-Id": user_a["id"]})
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(self.other.get("/api/orders/history").json()["total"], 1)
        self.assertEqual(self.other.get(f"/api/orders/history/{order_b}/file").content, exported_b.content)

    def test_saved_order_deletion_survives_new_process_and_login(self):
        user = self.register()
        removed = self.exported()
        retained = self.exported(quantity=3)
        removed_id = removed.headers["x-order-id"]
        retained_id = retained.headers["x-order-id"]
        response = self.client.delete(f"/api/orders/history/{removed_id}")
        self.assertEqual(response.status_code, 200, response.text)
        code = """
import hashlib,json,sys
from fastapi.testclient import TestClient
from app import app
data=json.load(sys.stdin)
with TestClient(app) as client:
    assert client.post('/api/accounts/login',json={'email':data['email'],'password':data['password']}).status_code==200
    history=client.get('/api/orders/history').json()
    removed=client.get('/api/orders/history/'+data['removed_id']+'/file')
    retained=client.get('/api/orders/history/'+data['retained_id']+'/file')
    assert retained.status_code==200
    print(json.dumps({'ids':[order['id'] for order in history['orders']],'total':history['total'],
                      'removed_status':removed.status_code,'retained_hash':hashlib.sha256(retained.content).hexdigest()}))
"""
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                                input=json.dumps({"email": user["email"], "password": PASSWORD,
                                                  "removed_id": removed_id, "retained_id": retained_id}),
                                text=True, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {"ids": [retained_id], "total": 1, "removed_status": 404,
                                                   "retained_hash": hashlib.sha256(retained.content).hexdigest()})

    def test_failed_storage_does_not_report_order_deleted(self):
        self.register()
        exported = self.exported()
        order_id = exported.headers["x-order-id"]
        with patch("account_store.delete_order", side_effect=OSError("private storage detail")):
            response = self.client.delete(f"/api/orders/history/{order_id}")
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("private storage detail", response.text)
        self.assertEqual(self.client.get("/api/orders/history").json()["total"], 1)
        self.assertEqual(self.client.get(f"/api/orders/history/{order_id}/file").content, exported.content)

    def test_successful_smtp_history_contains_exact_sent_attachment(self):
        self.register()
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            connection = smtp.return_value.__enter__.return_value
            connection.send_message.return_value = {}
            response = self.client.post("/api/orders/send", json=self.order())
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(set(response.json()), {"ok", "order_id", "order_saved"})
            self.assertIs(response.json()["order_saved"], True)
            attached = list(connection.send_message.call_args.args[0].iter_attachments())[0].get_payload(decode=True)
        order_id = response.json()["order_id"]
        history = self.client.get("/api/orders/history").json()["orders"]
        self.assertEqual([order["id"] for order in history], [order_id])
        self.assertEqual(history[0]["status"], "sent")
        self.assertEqual(self.client.get(f"/api/orders/history/{order_id}/file").content, attached)

    def test_anonymous_smtp_success_reports_order_not_saved_when_history_disabled(self):
        settings = {"ORDERS_HISTORY_DIR": "", "SMTP_HOST": "smtp.example.com",
                    "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            connection = smtp.return_value.__enter__.return_value
            connection.send_message.return_value = {}
            response = self.client.post("/api/orders/send", json=self.order())
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(set(response.json()), {"ok", "order_id", "order_saved"})
            self.assertIs(response.json()["ok"], True)
            self.assertIs(response.json()["order_saved"], False)
            self.assertTrue(response.json()["order_id"])
            attached = list(connection.send_message.call_args.args[0].iter_attachments())[0].get_payload(decode=True)
            self.assertEqual(load_workbook(io.BytesIO(attached))["Заказ"]["B3"].value, 2)
            self.assertEqual(self.client.get("/api/orders/history").status_code, 503)

    def test_smtp_is_not_attempted_when_initial_order_persistence_fails(self):
        self.register()
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("account_store.save_order", side_effect=OSError("storage unavailable")), \
                patch("app.smtplib.SMTP", autospec=True) as smtp:
            response = self.client.post("/api/orders/send", json=self.order())
            self.assertEqual(response.status_code, 503, response.text)
            self.assertNotIn("order_saved", response.json())
            smtp.assert_not_called()
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"], [])

    def test_failed_smtp_history_keeps_downloadable_order_and_failed_status(self):
        self.register()
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "sender@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", side_effect=smtplib.SMTPException("failure")):
            response = self.client.post("/api/orders/send", json=self.order())
        self.assertEqual(response.status_code, 502)
        history = self.client.get("/api/orders/history").json()["orders"]
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["status"], "failed")
        document = self.client.get(f"/api/orders/history/{history[0]['id']}/file")
        self.assertEqual(document.status_code, 200)
        self.assertEqual(load_workbook(io.BytesIO(document.content))["Заказ"]["B3"].value, 2)

    def test_unconfigured_smtp_is_failed_history_not_fake_sent_success(self):
        self.register()
        self.assertEqual(self.client.post("/api/orders/send", json=self.order()).status_code, 503)
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"][0]["status"], "failed")

    def test_unavailable_storage_never_returns_fake_saved_excel(self):
        self.register()
        with patch("account_store.save_order", side_effect=OSError("storage unavailable")):
            response = self.client.post("/api/orders/export", json=self.order())
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.client.get("/api/orders/history").json()["orders"], [])


class PostgresTransportTests(unittest.TestCase):
    def test_external_render_requires_tls_and_preserves_explicit_verification(self):
        external = "postgresql://test-user@dpg-test.oregon-postgres.render.com/catalog"
        self.assertEqual(postgres_tls_options(external), {"sslmode": "require"})
        self.assertEqual(postgres_tls_options(external + "?sslmode=verify-full"), {})
        self.assertEqual(postgres_tls_options("postgresql://test-user@127.0.0.1/catalog"), {})


if __name__ == "__main__":
    unittest.main()
