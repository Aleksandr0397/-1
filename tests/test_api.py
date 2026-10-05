import base64
import io
import json
import os
import smtplib
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import zipfile
import zlib
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
from email import policy
from email.parser import BytesParser
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import app


def price_file(rows=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Прайс Окунев, октябрь 2026"])
    sheet.append([])
    sheet.append(["Артикул", "Наименование товара", "Ед. изм.", "Цена, руб."])
    for row in rows or [
        ["A-01", "Кабель медный Ёлка", "м", "12,50"],
        ["A-02", "Лампа белая", "шт", 21.75],
        ["", "Без цены", "шт", "по запросу"],
    ]:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.data = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {
            "DATA_DIR": self.data.name,
            "INITIAL_CATALOG_FILE": "",
            "ORDERS_DATABASE_URL": "", "ORDERS_HISTORY_DIR": "",
            "ADMIN_TOKEN": "test-owner-password",
            "SMTP_HOST": "", "SMTP_FROM": "", "SMTP_USER": "",
            "SMTP_PASSWORD": "", "SMTP_PORT": "587",
            "SMTP_SECURITY": "starttls", "ORDERS_TO": "",
        })
        self.environment.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.environment.stop()
        self.data.cleanup()

    def upload(self, content=None, filename="price.xlsx", password="test-owner-password"):
        return self.client.post("/api/catalog/import", headers={"X-Admin-Token": password},
                                files={"file": (filename, content or price_file())})

    def order(self, catalog=None):
        catalog = catalog or self.upload().json()
        return {
            "catalog_version": catalog["catalog_version"],
            "customer": {"name": "Покупатель", "contact": "buyer@example.com", "comment": "Доставка"},
            "items": [{"product_id": catalog["products"][0]["id"], "quantity": "2"}],
        }

    def test_empty_catalog_and_runtime_config(self):
        self.assertEqual(self.client.get("/health").json(), {"status": "ok"})
        self.assertEqual(self.client.get("/api/catalog").json()["products"], [])
        config = self.client.get("/api/config").json()
        self.assertTrue(config["upload_enabled"])
        self.assertTrue(config["upload_requires_password"])
        self.assertFalse(config["mail_ready"])
        self.assertFalse(config["history_ready"])
        self.assertEqual(config["max_upload_mb"], 10)

    def test_import_detects_title_and_cyrillic_headers_and_persists(self):
        response = self.upload()
        self.assertEqual(response.status_code, 200, response.text)
        catalog = response.json()
        self.assertEqual(catalog["imported_count"], 2)
        self.assertEqual(catalog["skipped_count"], 1)
        self.assertEqual(catalog["products"][0]["name"], "Кабель медный Ёлка")
        self.assertEqual(catalog["products"][0]["price"], "12.50")
        self.assertEqual(catalog["products"][0]["unit"], "м")
        with TestClient(app) as fresh_client:
            stored = fresh_client.get("/api/catalog").json()
        self.assertEqual(stored["products"], catalog["products"])
        self.assertEqual(stored["catalog_version"], catalog["catalog_version"])
        self.assertTrue(stored["uploaded_at"])

    def initial_snapshot(self):
        snapshot = {"catalog_version": "published-price-v1", "filename": "Прайс.xlsx",
                    "uploaded_at": "2026-10-04T08:00:00+00:00",
                    "products": [{"id": "published-price-v1:1", "sku": "A-1", "name": "Товар из прайса",
                                  "unit": "шт", "price": "283.50"}]}
        path = os.path.join(self.data.name, "initial-catalog.json")
        with open(path, "w", encoding="utf-8") as stream:
            json.dump(snapshot, stream, ensure_ascii=False)
        return path, snapshot

    def test_fresh_database_bootstraps_snapshot_and_retains_version_and_ids(self):
        path, snapshot = self.initial_snapshot()
        with patch.dict(os.environ, {"INITIAL_CATALOG_FILE": path}):
            response = self.client.get("/api/catalog")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json(), snapshot)
            order = self.order(snapshot)
            exported = self.client.post("/api/orders/export", json=order)
            self.assertEqual(exported.status_code, 200, exported.text)
            sheet = load_workbook(io.BytesIO(exported.content)).active
            self.assertEqual(sheet["A3"].value, "Товар из прайса")
            self.assertEqual(sheet["E3"].value, 567)
        # Once persisted, the bundled snapshot is no longer needed to read the catalog.
        os.remove(path)
        self.assertEqual(self.client.get("/api/catalog").json(), snapshot)

    def test_bootstrap_never_overwrites_owner_import_even_after_restart(self):
        path, snapshot = self.initial_snapshot()
        with patch.dict(os.environ, {"INITIAL_CATALOG_FILE": path}):
            self.assertEqual(self.client.get("/api/catalog").json(), snapshot)
            imported = self.upload().json()
            self.assertNotEqual(imported["catalog_version"], snapshot["catalog_version"])
            with TestClient(app) as restarted:
                stored = restarted.get("/api/catalog").json()
            self.assertEqual(stored["products"], imported["products"])
            self.assertEqual(stored["catalog_version"], imported["catalog_version"])
            # An existing catalog bypasses even an absent snapshot file.
            os.remove(path)
            self.assertEqual(self.client.get("/api/catalog").json(), stored)

    def test_bootstrap_does_not_replace_catalog_imported_before_configuration(self):
        imported = self.upload().json()
        path, _ = self.initial_snapshot()
        with patch.dict(os.environ, {"INITIAL_CATALOG_FILE": path}):
            stored = self.client.get("/api/catalog").json()
        self.assertEqual(stored["catalog_version"], imported["catalog_version"])
        self.assertEqual(stored["products"], imported["products"])

    def test_invalid_bootstrap_is_atomic_and_can_be_repaired(self):
        path, snapshot = self.initial_snapshot()
        invalid = json.loads(json.dumps(snapshot))
        invalid["products"].append({**snapshot["products"][0], "id": "published-price-v1:2", "price": "NaN"})
        with open(path, "w", encoding="utf-8") as stream:
            json.dump(invalid, stream)
        with patch.dict(os.environ, {"INITIAL_CATALOG_FILE": path}):
            with self.assertRaises(RuntimeError):
                self.client.get("/api/catalog")
            with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM products").fetchone()[0], 0)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM catalog_meta").fetchone()[0], 0)
            with open(path, "w", encoding="utf-8") as stream:
                json.dump(snapshot, stream, ensure_ascii=False)
            self.assertEqual(self.client.get("/api/catalog").json(), snapshot)

    def test_concurrent_first_requests_seed_once_without_duplicate_products(self):
        path, snapshot = self.initial_snapshot()
        def fetch_catalog(_):
            with TestClient(app) as client:
                response = client.get("/api/catalog")
                self.assertEqual(response.status_code, 200, response.text)
                return response.json()
        with patch.dict(os.environ, {"INITIAL_CATALOG_FILE": path}), ThreadPoolExecutor(max_workers=4) as executor:
            responses = list(executor.map(fetch_catalog, range(4)))
        self.assertEqual(responses, [snapshot] * 4)

    def test_import_requires_configured_owner_secret(self):
        self.assertEqual(self.upload(password="wrong").status_code, 401)
        self.assertEqual(self.client.get("/api/catalog").json()["products"], [])
        with patch.dict(os.environ, {"ADMIN_TOKEN": ""}):
            self.assertEqual(self.upload().status_code, 503)
            self.assertFalse(self.client.get("/api/config").json()["upload_enabled"])

    def test_bad_file_does_not_replace_existing_catalog(self):
        catalog = self.upload().json()
        response = self.upload(b"This is not Excel")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.get("/api/catalog").json()["catalog_version"], catalog["catalog_version"])
        self.assertEqual(self.upload(filename="price.csv").status_code, 400)

    def test_catalog_version_endpoint_returns_only_public_current_or_archived_snapshot(self):
        self.assertEqual(self.client.get("/api/catalog/versions/unknown-version").status_code, 404)
        self.upload()
        first = self.client.get("/api/catalog").json()
        current = self.client.get("/api/catalog/versions/" + first["catalog_version"])
        self.assertEqual(current.status_code, 200, current.text)
        self.assertEqual(current.json(), first)
        self.upload(price_file([["NEW", "Товар нового прайса", "шт", "99,50"]]), filename="new-price.xlsx")
        second = self.client.get("/api/catalog").json()
        archived = self.client.get("/api/catalog/versions/" + first["catalog_version"])
        self.assertEqual(archived.status_code, 200, archived.text)
        self.assertEqual(archived.json(), first)
        self.assertEqual(set(archived.json()), {"catalog_version", "filename", "uploaded_at", "products"})
        for product in archived.json()["products"]:
            self.assertEqual(set(product), {"id", "sku", "name", "unit", "price"})
        self.assertEqual(self.client.get("/api/catalog/versions/" + second["catalog_version"]).json(), second)
        self.assertEqual(self.client.get("/api/catalog").json()["products"][0]["name"], "Товар нового прайса")
        self.assertEqual(self.client.get("/api/catalog/versions/still-unknown").status_code, 404)

    def test_multiple_price_replacements_preserve_exact_ids_names_prices_and_order_for_recovery(self):
        snapshots = []
        for index in range(3):
            self.assertEqual(self.upload(price_file([["SKU", f"Товар {index}", "шт", 10 + index],
                                                     ["SECOND", f"Вторая позиция {index}", "м", 20 + index]])).status_code, 200)
            snapshots.append(self.client.get("/api/catalog").json())
        for snapshot in snapshots:
            self.assertEqual(self.client.get("/api/catalog/versions/" + snapshot["catalog_version"]).json(), snapshot)
        self.assertEqual(self.client.get("/api/catalog").json(), snapshots[-1])
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            archived_versions = {row[0] for row in connection.execute("SELECT version FROM catalog_archive")}
            self.assertEqual(archived_versions, {snapshot["catalog_version"] for snapshot in snapshots[:-1]})
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM products").fetchone()[0], 2)

    def test_invalid_price_import_preserves_current_catalog_and_existing_archive(self):
        self.upload()
        first = self.client.get("/api/catalog").json()
        self.upload(price_file([["NEW", "Актуальный товар", "шт", 50]]))
        current = self.client.get("/api/catalog").json()
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            before = connection.execute("SELECT * FROM catalog_archive").fetchall()
        for invalid in (b"Not an Excel file", price_file([["BAD", "Без корректной цены", "шт", "по запросу"]])):
            self.assertEqual(self.upload(invalid).status_code, 400)
        self.assertEqual(self.client.get("/api/catalog").json(), current)
        self.assertEqual(self.client.get("/api/catalog/versions/" + first["catalog_version"]).json(), first)
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            self.assertEqual(connection.execute("SELECT * FROM catalog_archive").fetchall(), before)

    def test_archive_and_replacement_roll_back_together_if_new_product_write_fails(self):
        self.upload()
        first = self.client.get("/api/catalog").json()
        self.upload(price_file([["SECOND", "Второй прайс", "шт", 50]]))
        current = self.client.get("/api/catalog").json()
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            before = connection.execute("SELECT * FROM catalog_archive").fetchall()
            connection.execute("""CREATE TRIGGER reject_replacement BEFORE INSERT ON products
                                  BEGIN SELECT RAISE(ABORT, 'injected product write failure'); END""")
        with self.assertRaises(sqlite3.IntegrityError):
            self.upload(price_file([["THIRD", "Третий прайс", "шт", 75]]))
        self.assertEqual(self.client.get("/api/catalog").json(), current)
        self.assertEqual(self.client.get("/api/catalog/versions/" + first["catalog_version"]).json(), first)
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            self.assertEqual(connection.execute("SELECT * FROM catalog_archive").fetchall(), before)

    def test_catalog_archive_survives_new_process(self):
        self.upload()
        first = self.client.get("/api/catalog").json()
        self.upload(price_file([["NEW", "Новый прайс", "шт", 99]]))
        current = self.client.get("/api/catalog").json()
        code = """
import json,sys
from fastapi.testclient import TestClient
from app import app
data=json.load(sys.stdin)
with TestClient(app) as client:
    archived=client.get('/api/catalog/versions/'+data['old_version'])
    assert archived.status_code==200
    print(json.dumps({'archived':archived.json(),'current':client.get('/api/catalog').json()}))
"""
        result = subprocess.run([sys.executable, "-c", code], cwd=os.path.dirname(os.path.dirname(__file__)),
                                input=json.dumps({"old_version": first["catalog_version"]}),
                                text=True, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {"archived": first, "current": current})

    def test_catalog_archive_does_not_change_saved_order_or_accept_old_checkout(self):
        history_dir = os.path.join(self.data.name, "order-history")
        with patch.dict(os.environ, {"ORDERS_HISTORY_DIR": history_dir}):
            registered = self.client.post("/api/accounts/register", json={"email": "archive-buyer@example.com",
                                                                           "password": "archive-test-long-password"})
            self.assertEqual(registered.status_code, 200, registered.text)
            self.upload()
            first = self.client.get("/api/catalog").json()
            old_order = self.order(first)
            exported = self.client.post("/api/orders/export", json=old_order)
            self.assertEqual(exported.status_code, 200, exported.text)
            history = self.client.get("/api/orders/history").json()
            self.upload(price_file([["NEW", "Новый прайс", "шт", 99]]))
            self.assertEqual(self.client.get("/api/catalog/versions/" + first["catalog_version"]).json(), first)
            self.assertEqual(self.client.get("/api/orders/history").json(), history)
            self.assertEqual(self.client.get("/api/orders/history/" + exported.headers["x-order-id"] + "/file").content,
                             exported.content)
            self.assertEqual(self.client.post("/api/orders/export", json=old_order).status_code, 409)

    def test_concurrent_price_uploads_archive_every_superseded_public_snapshot(self):
        self.upload()
        first = self.client.get("/api/catalog").json()
        def upload_price(index):
            with TestClient(app) as client:
                response = client.post("/api/catalog/import", headers={"X-Admin-Token": "test-owner-password"},
                                       files={"file": (f"price-{index}.xlsx", price_file([[str(index), f"Товар {index}", "шт", index]]))})
                self.assertEqual(response.status_code, 200, response.text)
                return {field: response.json()[field] for field in ("catalog_version", "filename", "uploaded_at", "products")}
        with ThreadPoolExecutor(max_workers=2) as workers:
            uploaded = list(workers.map(upload_price, (1, 2)))
        for snapshot in (first, *uploaded):
            self.assertEqual(self.client.get("/api/catalog/versions/" + snapshot["catalog_version"]).json(), snapshot)
        with sqlite3.connect(os.path.join(self.data.name, "catalog.sqlite3")) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM catalog_archive").fetchone()[0], 2)

    def test_oversized_body_and_untrusted_body_rejected_before_parsing(self):
        body = b"x" * (10 * 1024 * 1024 + 256 * 1024 + 1)
        response = self.client.post("/api/catalog/import", content=body,
                                    headers={"X-Admin-Token": "test-owner-password", "Content-Type": "multipart/form-data; boundary=sample"})
        self.assertEqual(response.status_code, 413, response.text)
        response = self.client.post("/api/catalog/import", content=b"malformed multipart",
                                    headers={"X-Admin-Token": "wrong", "Content-Type": "multipart/form-data; boundary=sample"})
        self.assertEqual(response.status_code, 401, response.text)
        def chunked_body():
            yield b'--sample\r\nContent-Disposition: form-data; name="file"; filename="price.xlsx"\r\n\r\n'
            for _ in range(11):
                yield b"x" * 1024 * 1024
            yield b"\r\n--sample--\r\n"
        response = self.client.post("/api/catalog/import", content=chunked_body(),
                                    headers={"X-Admin-Token": "test-owner-password", "Content-Type": "multipart/form-data; boundary=sample"})
        self.assertEqual(response.status_code, 413, response.text)

    def test_legacy_xls_import(self):
        # Generated two-row BIFF fixture, compressed to keep source small; no xlwt dependency.
        fixture = (
            "eJztWE1IVFEU/u7zzDgz+DNjWmhko9BUNoRpQ1M4zmiULUqsFkkEpekiFI2hTS3KslkGgauijeCmjdWmH2pRuxaBUYtACLSWrYKCFjqvc868MS0DZ6H98L7HPfecc+8959zfd997PRWaGX9QM4uf0IoiZG0/vIt0hpM/LwTB5bYtbD73cbJd/FPw+3givR48LX1VLHMo8z0LC/fpBVPgA6dTOI/O4aH+8BqiXWPoMRJDgqnBHdaUoVqjqlB6Vuk6pfe05jOlKdXcUJrgujPmJKaSnQ1xZxV3W/VaVgax+0jbTKtmF6rwUlbxlZsmV9eDtvS5nsG/s6CWSjABnreO/qH+dM/gDCp5Aifw1Q4DX/I79XnY1a+t3oD135bqi5fRj1kEjMA+ows8w0u1gaTEYzZTihppL7XRHK7qwSqpBGYTq+PUQjFKUCs1sZRgOUZemA0s7OO0h8bqfDBVtJt2ssM4K1qYI5iINmskD0wd7WdtCwt+mHrHVIrC6MZhHJeNcYjauSKa4s3RWICd506K4JKTolR3UAnTPpQrH9J9FORw5+5+fnOktyt5WjUj2oHcO2SLdBs2d4xbcOMyLbGUskc0KL9D6TW1ulH5GqWVXIfzSFeVwxwc1TrXtTTCfpoVb5NbF/HbmM98Ovq4NvMxuZ35yY7ZS5WT75LjqOd3Wh+3l2cUURM1t28JniTzuXHOm/dKq385e3xW0Indzs2sKcc8AsqGlOYkGR2zIFnOWOWkIpaKFiRiiRzLZhnLRi3LWBnLq5GHHL1Y9Tp2jFqVybpsiUSOx1arAg91uFP4gQBcuHDhwoULFy5c/P+Qm6RcBeXuKTdOuVPK/VFujfJfZ55T9k//pHCxajiGYX4u8HfiAQxxnsbFgtbPev5YzdsyK2yT/18oOMHe0xhAr8YxUJBvAX+ZmcX9WXHDYMGufoeC/WcLiXOV/X8HuHnL8A=="
        )
        response = self.upload(zlib.decompress(base64.b64decode(fixture)), filename="price.xls")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["imported_count"], 1)
        self.assertEqual(response.json()["products"][0]["name"], "Товар XLS")
        self.assertEqual(response.json()["products"][0]["price"], "283.50")

    def test_export_selected_positions_decimal_total_and_customer(self):
        catalog = self.upload().json()
        order = self.order(catalog)
        response = self.client.post("/api/orders/export", json=order)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(".xlsx", response.headers["content-disposition"])
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        sheet = workbook.active
        rows = list(sheet.values)
        names = [cell for row in rows for cell in row]
        self.assertIn("Кабель медный Ёлка", names)
        self.assertNotIn("Лампа белая", names)
        self.assertEqual(list(sheet.values)[1], ("Наименование", "Заказ↓", "Ед. изм.", "Цена", "Сумма"))
        self.assertEqual(workbook["Информация"]["B3"].value, "Покупатель")
        product_row = next(row for row in rows if "Кабель медный Ёлка" in row)
        self.assertEqual(product_row[1:5], (2, "м", 12.5, 25))
        self.assertEqual(next(row for row in rows if "Итого" in row)[4], 25)
        self.assertEqual(sheet["B3"].number_format, "0")
        self.assertTrue(sheet.freeze_panes)

    def test_untrusted_strings_never_become_spreadsheet_formulas(self):
        catalog = self.upload(price_file([["=1+1", "=HYPERLINK(\"http://example.com\")", "+unit", 2]])).json()
        order = self.order(catalog)
        order["customer"] = {"name": "=2+2", "contact": "@SUM(1)", "comment": "-1+1"}
        response = self.client.post("/api/orders/export", json=order)
        self.assertEqual(response.status_code, 200, response.text)
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        values = [cell.value for sheet in workbook for row in sheet for cell in row]
        self.assertIn("=2+2", values)
        self.assertIn("=HYPERLINK(\"http://example.com\")", values)
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))

    def test_supplier_template_without_declared_dimensions(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([datetime(2026, 8, 8)])
        sheet.append(["Наименование", "Заказ↓", "Ед. изм.", "Цена", "Сумма"])
        sheet.append(["  Товар поставщика", None, "шт", "283,5", 0])
        original = io.BytesIO()
        workbook.save(original)
        updated = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(updated, "w") as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    import re
                    content = re.sub(rb'<dimension[^>]*/>', b'', content)
                target.writestr(entry, content)
        response = self.upload(updated.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["products"][0]["name"], "Товар поставщика")
        self.assertEqual(response.json()["products"][0]["price"], "283.50")

    def test_invalid_quantity_unknown_duplicate_and_empty_items(self):
        order = self.order()
        for quantity in [0, -1, "NaN", "Infinity", "1000000", True, None, "", "one", "1..0", {"bad": 1}]:
            with self.subTest(quantity=quantity):
                order["items"][0]["quantity"] = quantity
                response = self.client.post("/api/orders/export", json=order)
                self.assertEqual(response.status_code, 422, response.text)
        order["items"][0]["quantity"] = "2"
        order["items"][0]["product_id"] = "unknown"
        self.assertEqual(self.client.post("/api/orders/export", json=order).status_code, 422)
        order = self.order(self.client.get("/api/catalog").json())
        order["items"].append(dict(order["items"][0]))
        self.assertEqual(self.client.post("/api/orders/export", json=order).status_code, 422)
        order["items"] = []
        self.assertEqual(self.client.post("/api/orders/export", json=order).status_code, 422)

    def test_fractional_quantity_rejected_for_export_and_email(self):
        order = self.order()
        for quantity in [1.5, 0.001, "1.5", "0.001", "0.0001", "2,5", "1.999"]:
            for endpoint in ["/api/orders/export", "/api/orders/send"]:
                with self.subTest(quantity=quantity, endpoint=endpoint):
                    order["items"][0]["quantity"] = quantity
                    response = self.client.post(endpoint, json=order)
                    self.assertEqual(response.status_code, 422, f"Unexpected response status: {response.status_code}")

    def test_whole_quantity_and_integral_representations_export_as_integers(self):
        order = self.order()
        for quantity in [1, 2, "3", 1.0, "2.0", 999999]:
            with self.subTest(quantity=quantity):
                order["items"][0]["quantity"] = quantity
                response = self.client.post("/api/orders/export", json=order)
                self.assertEqual(response.status_code, 200, response.text)
                sheet = load_workbook(io.BytesIO(response.content)).active
                self.assertEqual(sheet["B3"].value, int(float(quantity)))
                self.assertIsInstance(sheet["B3"].value, int)
                self.assertEqual(sheet["B3"].number_format, "0")
                self.assertEqual(sheet["D3"].number_format, "#,##0.00")
    def test_reimport_invalidates_stale_basket(self):
        order = self.order()
        self.upload(price_file([["B", "Другой товар", "шт", 99]]))
        for endpoint in ["/api/orders/export", "/api/orders/send"]:
            response = self.client.post(endpoint, json=order)
            self.assertEqual(response.status_code, 409, response.text)

    def test_decimal_half_up_and_excel_precision_boundary(self):
        catalog = self.upload(price_file([["A", "Точное округление", "шт", "1.005"]])).json()
        order = self.order(catalog)
        response = self.client.post("/api/orders/export", json=order)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet["D3"].value, 1.01)
        self.assertEqual(sheet["E3"].value, 2.02)
        self.assertEqual(sheet["E4"].value, 2.02)
        catalog = self.upload(price_file([["A", "Очень дорогой товар", "шт", "999999998.99"]])).json()
        order = self.order(catalog)
        order["items"][0]["quantity"] = "999999"
        response = self.client.post("/api/orders/export", json=order)
        self.assertEqual(response.status_code, 422, response.text)

    def test_unconfigured_email_is_explicit_failure(self):
        order = self.order()
        response = self.client.post("/api/orders/send", json=order)
        self.assertEqual(response.status_code, 503, response.text)

    def test_customer_fields_are_optional_for_export_and_send(self):
        order = self.order()
        order["customer"] = {}
        response = self.client.post("/api/orders/export", json=order)
        self.assertEqual(response.status_code, 200, response.text)
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "orders@example.com", "ORDERS_TO": "owner@example.com"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            connection = smtp.return_value.__enter__.return_value
            connection.send_message.return_value = {}
            response = self.client.post("/api/orders/send", json=order)
            self.assertEqual(response.status_code, 200, response.text)
            message = connection.send_message.call_args.args[0]
            self.assertIn("Не указано", message.get_body(preferencelist=("plain",)).get_content())
            self.assertIsNone(message["Reply-To"])

    def test_email_has_actual_selected_xlsx_attachment(self):
        order = self.order()
        settings = {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "orders@example.com",
                    "ORDERS_TO": "owner@example.com", "SMTP_USER": "user", "SMTP_PASSWORD": "secret"}
        with patch.dict(os.environ, settings), patch("app.smtplib.SMTP", autospec=True) as smtp:
            connection = smtp.return_value.__enter__.return_value
            connection.send_message.return_value = {}
            response = self.client.post("/api/orders/send", json=order)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertTrue(response.json()["ok"])
            self.assertTrue(response.json()["order_id"])
            smtp.assert_called_once_with("smtp.example.com", 587, timeout=20)
            connection.starttls.assert_called_once()
            connection.login.assert_called_once_with("user", "secret")
            message = connection.send_message.call_args.args[0]
        parsed = BytesParser(policy=policy.default).parsebytes(message.as_bytes())
        self.assertEqual(parsed["To"], "owner@example.com")
        attachments = list(parsed.iter_attachments())
        self.assertEqual(len(attachments), 1)
        self.assertTrue(attachments[0].get_filename().endswith(".xlsx"))
        self.assertEqual(set(response.json()), {"ok", "order_id"})
        sheet = load_workbook(io.BytesIO(attachments[0].get_payload(decode=True))).active
        values = [cell.value for row in sheet for cell in row]
        self.assertIn("Кабель медный Ёлка", values)
        self.assertNotIn("Лампа белая", values)

    def test_smtp_failure_does_not_report_success(self):
        order = self.order()
        with patch.dict(os.environ, {"SMTP_HOST": "smtp.example.com", "SMTP_FROM": "orders@example.com", "ORDERS_TO": "owner@example.com"}), \
                patch("app.smtplib.SMTP", side_effect=smtplib.SMTPException("server failure")):
            response = self.client.post("/api/orders/send", json=order)
        self.assertEqual(response.status_code, 502, response.text)


if __name__ == "__main__":
    unittest.main()
